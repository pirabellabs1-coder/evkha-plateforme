"""Le lot du 31/10/2026 : l'inventaire des quatre études de test, corrigé d'un bloc.

Inventaire fait sans génération sur `03ba4f93`, `bc694571`, `d24977e6` et
`e30fce14` (67,28 € de tests le 30/10), puis validé par l'utilisateur. Chaque
classe a son test, et chacun échoue sur le code d'avant (règle 6), sauf les
contre-épreuves, nommées comme telles.

Fausses alertes — un contrôle qui frappe à tort est pire qu'absent (règle 2),
et ceux-ci sont GRAVES à la rédaction : ils font reprendre, et payer, un
chapitre juste.
- « 4 × 0,25 + … + 5 × 0,05 = 3,80 » lu « 5 × 0,05 = 3,80 » ;
- « Le chapitre 12 a classé dix risques » comparé au tableau du chapitre 13 ;
- 22 « fautes d'orthographe » sur 22 fausses : « commandes » / « commandés »,
  mots d'adresses web, « croit » / « croît » ;
- la règle des loisirs datés appliquée à la rétractation d'un professionnel.

Défauts réels du document :
- 108 espaces ordinaires entre un nombre et « € » ou « % » ;
- « 26 % des TPE-PME » devenu « 26 % des TPE », jamais renvoyé au rédacteur ;
- « AI Act » huit fois par étude ;
- des sources citées absentes de la liste, et l'inverse ;
- une couverture sans motif (Evangéline, 29/10/2026).
"""
from __future__ import annotations

from typing import Any

import pytest

from generation.orthographe import accent_final_seulement
from generation.relecture import Document, Reference, document_du_chapitre, relire
from generation.relecture.formules import _operations
from generation.relecture.redaction import _orthographe_a_deux_graphies


def _document(*paragraphes: str, chapitre: int = 17, tableau: Any = None) -> Document:
    blocs: list[dict[str, Any]] = [{"type": "paragraphe", "texte": p} for p in paragraphes]
    if tableau is not None:
        blocs.append({"type": "tableau", "tableau": tableau})
    return document_du_chapitre({"chapitre": chapitre, "titre": "Chapitre", "blocs": blocs})


# ── 1. Le calcul enchaîné ────────────────────────────────────────────────────

SOMME_PONDEREE = (
    "Exemple de calcul pour les villes moyennes : 4 × 0,25 + 3 × 0,15 + 3 × 0,10 "
    "+ 4 × 0,20 + 4 × 0,15 + 4 × 0,10 + 5 × 0,05 = 3,80."
)


def test_une_somme_ponderee_juste_n_est_pas_un_calcul_faux() -> None:
    """LE cas de `e30fce14`, chapitre 17. Sur le code d'avant : constat GRAVE."""
    assert _operations(_document(SOMME_PONDEREE)) == []


def test_une_somme_ponderee_fausse_est_signalee() -> None:
    """CONTRE-ÉPREUVE : l'expression entière est jugée, pas abandonnée."""
    (constat,) = _operations(_document(SOMME_PONDEREE.replace("= 3,80", "= 4,20")))

    assert "4 × 0,25" in constat.detail
    assert "4,20" in constat.detail


def test_une_note_sur_cinq_ne_se_juge_pas_a_un_demi_point_pres() -> None:
    """La tolérance suit l'écriture : « 3,80 » se juge au centième."""
    assert _operations(_document("Note : 2 × 1,50 + 0,40 = 3,40.")) == []
    assert len(_operations(_document("Note : 2 × 1,50 + 0,40 = 3,60."))) == 1


def test_une_operation_simple_fausse_reste_signalee() -> None:
    """CONTRE-ÉPREUVE : le cas simple que le contrôle savait juger."""
    assert len(_operations(_document("Le panier vaut 25 × 40 = 1 200 €."))) == 1


def test_une_egalite_en_chaine_juge_chaque_etape() -> None:
    """« a = b = c » : l'étape suivante repart APRÈS le résultat (relecture du
    31/10). Une deuxième étape fausse se voit, une juste ne crie pas."""
    assert _operations(_document("Soit 2 × 3 = 6 × 4 = 24.")) == []
    (constat,) = _operations(_document("Soit 2 × 3 = 6 × 4 = 30."))
    assert "6 × 4" in constat.detail


def test_la_tolerance_suit_la_derniere_unite_ecrite_a_son_echelle() -> None:
    """« 1,2 M€ » se juge à 100 000 € près ; « 0,66 » tronqué n'est pas faux.
    Une seule tolérance, celle de `valeurs.tolerance_ecrite` (règle 5)."""
    assert _operations(_document("Soit 400 000 € × 3,1 = 1,2 M€.")) == []
    assert _operations(_document("Le ratio vaut 2 ÷ 3 = 0,66.")) == []
    # CONTRE-ÉPREUVE : un écart de plusieurs unités reste faux.
    assert _operations(_document("Soit 400 000 € × 3,1 = 1,5 M€."))


def test_une_fourchette_n_est_pas_une_soustraction() -> None:
    """« 10-15 % » est une fourchette : la lire « 10 − 15 » serait un faux calcul."""
    assert _operations(_document("Une marge de 10-15 % = 12 % en moyenne.")) == []


@pytest.mark.parametrize(
    "phrase",
    ["Prix : 100 € − 10 % + 5 € = 95 €.", "Soit 45 € × 1 200 − 10 % = 48 600 €."],
)
def test_un_pourcentage_dans_une_chaine_s_applique_au_montant(phrase: str) -> None:
    assert _operations(_document(phrase)) == []


# ── 2. L'annonce qui parle d'un autre chapitre ───────────────────────────────


def _comptage(texte: str) -> list[Any]:
    tableau = {
        "entetes": ["Risque", "Probabilité"],
        "lignes": [[f"Risque {n}", "Moyenne"] for n in range(1, 8)],
    }
    document = _document(texte, chapitre=13, tableau=tableau)
    return [c for c in relire(document, Reference()) if c.classe == "comptage"]


def test_une_annonce_du_chapitre_12_ne_se_compte_pas_au_13() -> None:
    """LE cas de `e30fce14` : la carte place sept des dix risques du chapitre 12."""
    assert _comptage(
        "Le chapitre 12 a classé dix risques. Sept viennent de l'extérieur."
    ) == []


def test_une_annonce_fausse_du_chapitre_reste_signalee() -> None:
    """CONTRE-ÉPREUVE : sans renvoi à un autre chapitre, dix pour sept se voit."""
    assert _comptage("Voici les dix risques de la carte.")


def test_un_chapitre_cite_apres_l_annonce_ne_l_excuse_pas() -> None:
    """CONTRE-ÉPREUVE : seul un renvoi AVANT l'annonce en change l'objet."""
    assert _comptage("La carte place dix risques, détaillés au chapitre 12.")


# ── 3. Les fausses fautes d'orthographe ──────────────────────────────────────


@pytest.mark.parametrize("mot", ["commandés", "limités", "notés", "chiffrés"])
def test_un_participe_au_pluriel_a_son_jumeau_legitime(mot: str) -> None:
    assert accent_final_seulement(mot) is True


@pytest.mark.parametrize("mot", ["données", "précise", "études"])
def test_un_accent_qui_ne_se_dedouble_pas_reste_suspect(mot: str) -> None:
    """CONTRE-ÉPREUVE : « donnees » n'est pas un mot."""
    assert accent_final_seulement(mot) is False


def test_les_mots_d_une_adresse_ne_sont_pas_des_fautes() -> None:
    constats = _orthographe_a_deux_graphies(_document(
        "Le Baromètre France Num 2026 et la stratégie numérique.",
        "https://www.francenum.gouv.fr/ guides-et-conseils/strategie- numerique/"
        "barometre-france-num",
        "Selon le communiqué (presse.economie.gouv.fr), l'Économie suit.",
    ))

    assert constats == []


def test_croit_et_croit_accentue_sont_deux_mots() -> None:
    assert _orthographe_a_deux_graphies(_document(
        "Le client croit être autonome.", "Le marché croît de 23 % par an.",
    )) == []


def test_une_vraie_faute_d_accent_reste_signalee() -> None:
    """CONTRE-ÉPREUVE : « donnees » face à « données »."""
    assert _orthographe_a_deux_graphies(_document(
        "Les donnees de la douane couvrent 2024.", "Les données couvrent 2025.",
    ))


# ── 4. La rétractation d'un professionnel ────────────────────────────────────


def test_la_retractation_du_professionnel_n_est_pas_une_erreur_de_loisirs() -> None:
    """L221-3 cité avec L221-18 : un client professionnel, pas un loisir daté."""
    document = _document(
        "Nos ateliers de formation et nos soirées d'information sont ouverts.",
        "Un client de cinq salariés au plus qui signe sur son chantier a droit à "
        "14 jours de rétractation (articles L221-3 et L221-18 du code de la "
        "consommation).",
    )
    assert not [c for c in relire(document, Reference()) if "L221-28" in c.detail]


def test_la_retractation_d_un_atelier_date_reste_une_erreur() -> None:
    """CONTRE-ÉPREUVE : le cas pour lequel la règle existe."""
    document = _document(
        "Nos ateliers, soirées et week-ends sont réservés à date fixe.",
        "Le client dispose de 14 jours de rétractation (article L221-18).",
    )
    assert [c for c in relire(document, Reference()) if "L221-28" in c.detail]


# ── 5. L'insécable avant « € » et « % », dans le Word ───────────────────────


def test_le_word_ne_garde_aucune_espace_ordinaire_avant_l_unite() -> None:
    from docx import Document as DocumentWord

    from generation.rendu_word.depuis_json import poser_les_insecables

    document = DocumentWord()
    document.add_paragraph("Une croissance de 23,4 % par an et 30 000 € HT.")
    coupe = document.add_paragraph()
    coupe.add_run("soit 1 200")
    coupe.add_run(" € par mission")

    remplacees = poser_les_insecables(document)

    textes = [p.text for p in document.paragraphs]
    assert remplacees == 3
    assert textes[0] == "Une croissance de 23,4 % par an et 30 000 € HT."
    assert textes[1] == "soit 1 200 € par mission"
    assert poser_les_insecables(document) == 0, "idempotente"


def test_l_insecable_se_pose_a_travers_les_morceaux_et_en_pied_de_page() -> None:
    """Relecture du 31/10 : espace en fin de morceau, espace seule dans son
    morceau, espace fine sécable, pied de page."""
    from docx import Document as DocumentWord

    from generation.rendu_word.depuis_json import poser_les_insecables

    document = DocumentWord()
    fin = document.add_paragraph()
    fin.add_run("soit 30 000 ")
    fin.add_run("€ HT")
    seule = document.add_paragraph()
    for morceau in ("soit 30 000", " ", "€"):
        seule.add_run(morceau)
    document.add_paragraph("fine 30 000 €")
    document.sections[0].footer.paragraphs[0].text = "dont 12 % du total"

    assert poser_les_insecables(document) == 4
    assert [p.text for p in document.paragraphs] == [
        "soit 30 000 € HT", "soit 30 000 €", "fine 30 000 €",
    ]
    assert document.sections[0].footer.paragraphs[0].text == "dont 12 % du total"
    # CONTRE-ÉPREUVE : l'espace fine ENTRE les milliers ne bouge pas.
    assert "30 000" in document.paragraphs[2].text


def test_la_passe_part_avant_l_enregistrement_de_l_etude() -> None:
    """Le câble, pas seulement la fonction (règle 8)."""
    import inspect

    from generation.rendu_word import depuis_json

    source = inspect.getsource(depuis_json.rendre_etude)
    assert source.index("poser_les_insecables(document)") < source.index("document.save(")


# ── 6. Le glissement de population est réparable ─────────────────────────────


def test_le_glissement_fait_reecrire_son_chapitre() -> None:
    from generation.controle_final import REPARABLES_PAR_CHAPITRE
    from generation.verification.objet_des_chiffres import CONTROLE

    assert "glissement_de_population" in REPARABLES_PAR_CHAPITRE
    # CONTRE-ÉPREUVE : objets changés et répétitions restent en mesure seule.
    assert CONTROLE not in REPARABLES_PAR_CHAPITRE


# ── 7. « AI Act » ────────────────────────────────────────────────────────────


REGLEMENT = "règlement européen sur l'IA"


#: La réécriture automatique du texte a été RETIRÉE à la relecture du
#: 31/10/2026 : elle abîmait des phrases justes (« l'AI Act » → « le règlement
#: … » au milieu d'une élision, une parenthèse explicative effacée). Le nom
#: français se demande au rédacteur, dans la charte.


def test_la_charte_nomme_le_reglement_en_francais() -> None:
    from generation.chapitres.langage_simple import REGLEMENT_IA
    from generation.chapitres.runner import _SYSTEME

    assert REGLEMENT_IA == REGLEMENT
    assert f"« le {REGLEMENT} »" in _SYSTEME
    assert "forces, faiblesses, opportunités et menaces" in _SYSTEME


def _paiement_des_fournisseurs(phrase: str) -> list[Any]:
    def chapitre(numero: int, texte: str) -> list[Any]:
        return document_du_chapitre({
            "chapitre": numero, "titre": "C",
            "blocs": [{"type": "paragraphe", "texte": texte}],
        }).sections

    sections = chapitre(
        9, "Les salles de formation demandent un acompte de 30 % pour bloquer la date."
    ) + chapitre(14, phrase)
    return [
        c for c in relire(Document(sections=sections), Reference())
        if "ournisseur" in c.detail
    ]


def test_le_nom_du_reglement_n_est_pas_un_paiement() -> None:
    """Porte finale du 31/10 : la consigne de la charte déclenchait un
    constat GRAVE « politique de paiement contradictoire » (règle 5)."""
    from generation.chapitres.langage_simple import REGLEMENT_IA

    assert _paiement_des_fournisseurs(
        f"Le {REGLEMENT_IA} impose aux fournisseurs de modèles une documentation ; "
        "les obligations des utilisateurs s'appliquent après le 2 août 2026."
    ) == []


def test_des_prestataires_regles_apres_restent_une_version() -> None:
    """CONTRE-ÉPREUVE : le sens « paiement » de règlement compte toujours."""
    assert _paiement_des_fournisseurs(
        "Le règlement des prestataires intervient après la prestation."
    )


def test_le_texte_ecrit_n_est_plus_reecrit() -> None:
    """CONTRE-ÉPREUVE de la réparation retirée : le texte du rédacteur reste le sien."""
    from generation.chapitres.langage_simple import simplifier_texte

    phrase = "le règlement (UE) 2024/1689, appelé AI Act, qui s'applique"
    assert simplifier_texte(phrase, sigles_admis=True) == phrase


# ── 8. Le chapitre des sources reçoit les sources citées ────────────────────


@pytest.mark.django_db
def test_le_chapitre_des_sources_recoit_les_sources_citees() -> None:
    from decimal import Decimal

    from catalog.models import DeliverableType, Offer
    from customers.models import Customer
    from generation.chapitres.runner import _bloc_sources_citees
    from generation.models import ChapterGeneration, ChapterStatus, GenerationJob
    from orders.models import Order

    offre = Offer.objects.create(
        name="EM", slug="em-sources-citees", deliverable_type=DeliverableType.MARKET_STUDY,
    )
    commande = Order.objects.create(
        systeme_order_id="cmd-sources-citees",
        customer=Customer.objects.create(email="sources@exemple.fr"), offer=offre,
    )
    job = GenerationJob.objects.create(
        order=commande, deliverable_type=DeliverableType.MARKET_STUDY,
        budget_eur=Decimal("18"),
    )
    ChapterGeneration.objects.create(
        job=job, chapter_number=4, chapter_title="Usages", prompt_key="em.04.x",
        status=ChapterStatus.DONE,
        payload={"chapitre": 4, "titre": "Usages", "blocs": [
            {"type": "paragraphe",
             "texte": "Quatre TPE sur dix utilisent l'IA (Baromètre France Num, 2026)."},
            {"type": "paragraphe", "texte": "Source : Insee, 2025"},
        ]},
    )
    sources = ChapterGeneration.objects.create(
        job=job, chapter_number=21, chapter_title="Sources",
        prompt_key="em.21.sources_methodologie",
    )
    autre = ChapterGeneration.objects.create(
        job=job, chapter_number=15, chapter_title="Ressources",
        prompt_key="str.15.arbitrages_ressources",
    )

    bloc = _bloc_sources_citees(sources)

    assert "SOURCES CITÉES DANS L'ÉTUDE" in bloc
    assert "Insee (chapitre 4)" in bloc
    # Le chapitre 15 n'est pas écrit : ses sources manquent, le bloc le dit.
    assert "Chapitres pas encore écrits (15)" in bloc
    # CONTRE-ÉPREUVE : un chapitre qui n'est pas celui des sources ne reçoit rien.
    assert _bloc_sources_citees(autre) == ""


# ── 9. Des motifs sur la couverture ──────────────────────────────────────────


def test_la_couverture_porte_des_motifs_derriere_le_texte() -> None:
    from generation.rendu_word import composants
    from generation.rendu_word.gabarit import charger_gabarit
    from generation.rendu_word.palette import REF_CREME, REF_PRUNE, construire_palette

    document = charger_gabarit()
    palette = construire_palette(primaire=REF_PRUNE, fond_clair=REF_CREME)
    composants.couverture(document, palette, titre="Étude", client="Client")

    xml = document.element.xml
    assert xml.count("<v:oval") == 3
    assert palette.secondaire in xml
    # Derrière le texte, et positionnés sur la page sans alignement forcé à
    # gauche — sinon Word les empile sous le titre (aperçu du 31/10/2026).
    for forme in xml.split("<v:oval")[1:]:
        style = forme.split('style="', 1)[1].split('"', 1)[0]
        assert "z-index:-" in style
        assert "mso-position-horizontal:left" not in style


def test_la_quatrieme_de_couverture_reste_sobre() -> None:
    """CONTRE-ÉPREUVE : la demande portait sur la couverture."""
    from generation.rendu_word import composants
    from generation.rendu_word.gabarit import charger_gabarit
    from generation.rendu_word.palette import REF_CREME, REF_PRUNE, construire_palette

    document = charger_gabarit()
    palette = construire_palette(primaire=REF_PRUNE, fond_clair=REF_CREME)
    composants.quatrieme_couverture(document, palette, nom="EVKHA", mentions=[])

    assert "<v:oval" not in document.element.xml
