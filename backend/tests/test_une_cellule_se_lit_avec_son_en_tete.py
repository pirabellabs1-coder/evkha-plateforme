"""Une cellule de tableau se juge avec son en-tête et sa ligne, comme le lecteur la lit.

## Le défaut mesuré

Corpus du 14/09/2026 : 212 des 382 « chiffres hors socle » d'une étude
concurrentielle étaient des cellules du chapitre 6 jugées SEULES — « 250 000 € »,
« 0,029 % ». Or ce chapitre doit, par cahier des charges, estimer le chiffre
d'affaires et la part de chaque concurrent, et il le fait en tableau sous un
en-tête « CA estimé ». Le lecteur voit l'en-tête ; le contrôle ne le voyait pas,
et accusait d'invention une estimation déclarée (règle 2).

La correction ne crée aucune exception pour les tableaux : la cellule reçoit
pour « phrase » son en-tête et sa ligne, et les règles de la prose s'appliquent.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from generation.socle.referentiel import Fiabilite, Perimetre
from generation.socle.schema import DonneeSocle, Socle, Zone
from generation.verification.controles import controler_chiffres_hors_socle
from generation.verification.lecture import lire_livrable


def _socle() -> Socle:
    return Socle(
        secteur="conseil", zone=Zone(pays="France"), date_socle=date(2026, 9, 14),
        donnees=[DonneeSocle(
            id="marche_national_taille", libelle="Marché national", valeur=850,
            unite="MEUR", annee=2025, perimetre=Perimetre.NATIONAL,
            fiabilite=Fiabilite.OBSERVEE, source="Xerfi",
        )],
    )


def _docx(tmp_path: Path, entetes: list[str], lignes: list[list[str]]) -> Path:
    from docx import Document

    document = Document()
    document.add_paragraph("Estimation des concurrents.")
    table = document.add_table(rows=1 + len(lignes), cols=len(entetes))
    for colonne, entete in enumerate(entetes):
        table.cell(0, colonne).text = entete
    for rang, ligne in enumerate(lignes, start=1):
        for colonne, valeur in enumerate(ligne):
            table.cell(rang, colonne).text = valeur
    chemin = tmp_path / "etude.docx"
    document.save(str(chemin))
    return chemin


def _signales(chemin: Path) -> set[str]:
    anomalies = controler_chiffres_hors_socle(lire_livrable(chemin), _socle())
    return {a.detail.split(" »")[0].lstrip("« ") for a in anomalies}


def test_une_estimation_en_tableau_n_est_pas_une_invention(tmp_path: Path) -> None:
    chemin = _docx(
        tmp_path,
        ["Concurrent", "CA estimé", "Part de marché estimée"],
        [["Cabinet de Nantes", "250 000 €", "0,029 %"],
         ["ESN de Lille", "350 000 €", "0,041 %"]],
    )
    assert _signales(chemin) == set()


def test_un_chiffre_d_affaires_sans_estimation_declaree_reste_signale(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : l'en-tête n'excuse que ce qu'il déclare."""
    chemin = _docx(
        tmp_path,
        ["Concurrent", "Chiffre d'affaires", "Part de marché estimée"],
        [["Cabinet de Nantes", "250 000 €", "0,029 %"]],
    )
    assert "250 000 €" in _signales(chemin)
    assert "0,029 %" not in _signales(chemin)


def test_une_croissance_reste_un_fait_de_marche_meme_estimee(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : comme en prose, une croissance ne se « déduit » pas."""
    chemin = _docx(
        tmp_path,
        ["Concurrent", "Croissance estimée"],
        [["Cabinet de Nantes", "12 %"]],
    )
    assert "12 %" in _signales(chemin)


def test_le_motif_montre_l_en_tete_et_la_ligne(tmp_path: Path) -> None:
    """Règle 2 : « 250 000 € » seul est introuvable dans un document à dix tableaux."""
    chemin = _docx(
        tmp_path, ["Concurrent", "Chiffre d'affaires"], [["Cabinet de Nantes", "250 000 €"]],
    )
    anomalies = controler_chiffres_hors_socle(lire_livrable(chemin), _socle())
    assert any("Chiffre d'affaires" in a.extrait and "Cabinet de Nantes" in a.extrait
               for a in anomalies)


# ── Re-mesure après déploiement : quatre classes de plus ────────────────────


def test_le_commentaire_d_une_autre_colonne_n_annule_pas_l_estimation(tmp_path: Path) -> None:
    """« Croissance proche du marché » en dernière colonne ne qualifie pas le CA estimé."""
    chemin = _docx(
        tmp_path,
        ["Acteur", "CA 2025 (estimé)", "CA médian 2026 (estimé)", "Lecture"],
        [["France d'Or", "7,3 M€", "8,5 M€", "Croissance proche du marché"]],
    )
    assert _signales(chemin) == set()


def test_une_part_calculee_dans_sa_ligne_n_est_pas_inventee(tmp_path: Path) -> None:
    """200 000 € / 850 M€ du socle = 0,024 %."""
    chemin = _docx(
        tmp_path,
        ["Acteur", "CA estimé", "Part du marché national (%)"],
        [["Intégrateur IA", "200 000 €", "0,024 %"]],
    )
    assert "0,024 %" not in _signales(chemin)


def test_une_part_qui_ne_tombe_pas_juste_reste_signalee(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : 200 000 € / 850 M€ ne fait pas 3,7 %."""
    chemin = _docx(
        tmp_path,
        ["Acteur", "CA estimé", "Part du marché national (%)"],
        [["Intégrateur IA", "200 000 €", "3,7 %"]],
    )
    assert "3,7 %" in _signales(chemin)


def test_sans_en_tete_de_part_un_rapport_juste_ne_suffit_pas(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : un taux quelconque ne devient pas une part par coïncidence."""
    chemin = _docx(
        tmp_path,
        ["Acteur", "CA estimé", "Taux de marge"],
        [["Intégrateur IA", "200 000 €", "0,024 %"]],
    )
    assert "0,024 %" in _signales(chemin)


def test_une_ligne_qui_porte_son_adresse_est_sourcee(tmp_path: Path) -> None:
    chemin = _docx(
        tmp_path,
        ["Organisme", "Adresse", "Donnée retenue"],
        [["Service Public Entreprendre", "https://entreprendre.service-public.gouv.fr/F38497",
          "Seuil de régime simplifié : 83 600 €"]],
    )
    assert "83 600 €" not in _signales(chemin)


def test_un_seuil_legal_cite_par_son_article_est_source(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph(
        "Le plafond de la micro-entreprise est de 77 700 € pour les prestations de "
        "services (article 50-0 du code général des impôts)."
    )
    chemin = tmp_path / "loi.docx"
    document.save(str(chemin))
    assert "77 700 €" not in _signales(chemin)


def test_une_cellule_qui_est_une_phrase_se_juge_aussi_par_elle_meme(tmp_path: Path) -> None:
    """Stratégie `0ad5155b` : l'estimation déclarée DANS la cellule compte."""
    chemin = _docx(
        tmp_path,
        ["Critère", "Verdict", "Pourquoi"],
        [["Marché porteur", "Modérément",
          "le marché national est estimé à 150 M€ sans source dédiée, et l'objectif "
          "à trois ans ne représente que 0,46 % de ce total"]],
    )
    assert "0,46 %" not in _signales(chemin)
