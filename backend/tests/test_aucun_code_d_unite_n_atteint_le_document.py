"""Aucun code d'unité de stockage n'atteint le document livré.

29/09/2026, business plan ÉCLORE (107 pages), pages 103 et 106 : l'annexe
« D'où viennent les chiffres de cette étude » imprimait « 2 000 000 unite »,
« 30 000 MEUR », « 23 223,86 EUR ». `annexe_chiffres._valeur` recopiait le code
de stockage de l'unité, alors que `montant_lisible` existait et avait été
appliqué au tableau de repli le 26/09 : l'exemple corrigé, pas la classe
(règle 4).

Ce test vise la CLASSE : un socle qui porte une donnée dans CHAQUE famille
d'unité, un document rendu jusqu'au `.docx`, et aucun code — `EUR`, `MEUR`,
`unite`, `annees`, `note_sur_5`… — dans ce que le lecteur lit, ni dans les
chaînes écrites sur les figures. Il échoue sur le code d'avant (l'annexe), et
la contre-épreuve vérifie que les valeurs lisibles, elles, sont bien là.
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from generation.chapitres.schema import ChapitrePayload
from generation.rendu_word.assemblage import assembler_etude
from generation.rendu_word.depuis_json import rendre_etude
from generation.socle.referentiel import Fiabilite, Perimetre
from generation.socle.schema import (
    DonneeSocle,
    Socle,
    Zone,
    unite_lisible,
    unites_monetaires,
)

pytestmark = pytest.mark.django_db

#: Codes de stockage non monétaires : ceux que `unite_lisible` traduit.
_NON_MONETAIRES = ("unite", "millier", "million", "annees", "ratio", "note_sur_5", "note_sur_10")

#: Tout code qui ne doit JAMAIS paraître tel quel : ceux dont la forme lisible
#: diffère. Construit depuis le socle, pas recopié (règle 5).
_CODES_INTERDITS = sorted(
    code for code in (*unites_monetaires(), *_NON_MONETAIRES)
    if unite_lisible(code) != code
)
_CODE = re.compile(r"(?<![\w-])(?:" + "|".join(map(re.escape, _CODES_INTERDITS)) + r")(?![\w-])")


def _donnee(identifiant: str, libelle: str, valeur: float, unite: str) -> DonneeSocle:
    return DonneeSocle(
        id=identifiant, libelle=libelle, valeur=valeur, unite=unite, annee=2026,
        perimetre=Perimetre.ENTREPRISE, source="données du projet",
        fiabilite=Fiabilite.DECLAREE,
    )


def _socle() -> Socle:
    return Socle(
        secteur="bien-être en pleine nature",
        zone=Zone(pays="France"),
        date_socle=date(2026, 9, 29),
        donnees=[
            # Les trois valeurs relevées sur ÉCLORE, puis une par famille.
            _donnee("visiteurs_zone", "Visiteurs de la zone", 2_000_000, "unite"),
            _donnee("marche_national", "Marché national du bien-être", 30_000, "MEUR"),
            _donnee("resultat_net_an3", "Résultat net — exercice 3", 23_223.86, "EUR"),
            _donnee("budget_demarrage", "Budget de démarrage", 5, "kEUR"),
            _donnee("marche_mondial", "Marché mondial", 1.2, "MdEUR"),
            # Petit devant son échelle : s'écrivait « 0 Md€ » une fois le code
            # « MdEUR » traduit — un zéro faux (29/09/2026).
            _donnee("som", "Marché atteignable", 0.0003, "MdEUR"),
            _donnee("taux_remplissage", "Taux de remplissage", 62, "%"),
            _donnee("duree_amortissement", "Durée d'amortissement", 5, "annees"),
            _donnee("note_accueil", "Note d'accueil", 4.5, "note_sur_5"),
        ],
    )


def _chapitre() -> ChapitrePayload:
    ids = ["marche_national", "marche_mondial", "budget_demarrage"]
    return ChapitrePayload.model_validate({
        "chapitre": 1,
        "titre": "Le marché",
        "accroche": "Un marché large.",
        "blocs": [
            {"type": "paragraphe", "texte": "Le marché est large."},
            {"type": "graphique", "graphique": {
                "type": "entonnoir", "titre": "Du marché au projet", "donnees_ids": ids,
            }},
            # Refusée au dessin (unités mêlées) : imprimée en tableau de repli.
            {"type": "graphique", "graphique": {
                "type": "barres", "titre": "Repères mêlés",
                "donnees_ids": ["visiteurs_zone", "resultat_net_an3", "duree_amortissement"],
            }},
        ],
        "donnees_utilisees": [*ids, "taux_remplissage", "note_accueil"],
        "resume": "Résumé.",
    })


def _texte_du_docx(chemin: Path) -> str:
    document = Document(str(chemin))
    morceaux = [p.text for p in document.paragraphs]
    for table in document.tables:
        for ligne in table.rows:
            morceaux.extend(cellule.text for cellule in ligne.cells)
    for section in document.sections:
        morceaux.extend(p.text for p in section.header.paragraphs)
        morceaux.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(morceaux)


def _chaines(objet: Any) -> list[str]:
    if isinstance(objet, str):
        return [objet]
    if isinstance(objet, dict):
        return [c for valeur in objet.values() for c in _chaines(valeur)]
    if isinstance(objet, list | tuple):
        return [c for valeur in objet for c in _chaines(valeur)]
    return []


@pytest.fixture
def rendu(tmp_path: Path) -> tuple[dict[str, Any], str]:
    etude, _ = assembler_etude(
        socle=_socle(), chapitres=[_chapitre()], titre="Business plan",
        marque={"nom": "ÉCLORE"},
    )
    chemin = rendre_etude(etude, tmp_path / "unites.docx")
    return etude, _texte_du_docx(chemin)


def test_le_document_ne_porte_aucun_code_d_unite(rendu: tuple[dict[str, Any], str]) -> None:
    """Le défaut exact d'ÉCLORE, et toute sa classe."""
    _etude, texte = rendu
    trouves = sorted(set(_CODE.findall(texte)))
    assert not trouves, f"codes d'unité imprimés dans le document : {trouves}"


def test_les_figures_ne_portent_aucun_code_d_unite(rendu: tuple[dict[str, Any], str]) -> None:
    """Les chaînes dessinées DANS les images échappent à la lecture du texte."""
    etude, _texte = rendu
    figures = [
        bloc for section in [*etude["chapitres"], *etude.get("annexes", [])]
        for bloc in section["blocs"] if bloc["type"] == "graphique"
    ]
    assert figures, "le test doit traverser au moins une figure"
    trouves = sorted({m for bloc in figures for c in _chaines(bloc) for m in _CODE.findall(c)})
    assert not trouves, f"codes d'unité écrits sur les figures : {trouves}"


def test_les_valeurs_lisibles_sont_bien_imprimees(rendu: tuple[dict[str, Any], str]) -> None:
    """CONTRE-ÉPREUVE : la valeur ne disparaît pas avec son code."""
    _etude, texte = rendu
    compact = texte.replace(" ", " ").replace(" ", " ")
    # « 300 000,00 € » : l'annexe écrit ses montants en euros au même arrondi que
    # « 23 223,86 € » (30/09/2026, « arrondi identique pour tous les montants d'un
    # même tableau »).
    for attendu in ("30 000 M€", "23 223,86 €", "2 000 000", "1,2 Md€", "5 k€", "300 000,00 €"):
        assert attendu in compact, f"« {attendu} » absent du document"
    assert not re.search(r"(?m)^0 \w*€$", compact), "un montant réel écrit zéro"


def test_le_detecteur_attrape_le_texte_d_eclore() -> None:
    """Contre-épreuve du détecteur lui-même, sur le texte relevé le 29/09/2026."""
    releve = "Visiteurs | 2 000 000 unite | 2026\nMarché | 30 000 MEUR\n23 223,86 EUR"
    assert sorted(set(_CODE.findall(releve))) == ["EUR", "MEUR", "unite"]
    assert not _CODE.findall("Un marché de 30 000 M€, soit une unité de mesure.")
