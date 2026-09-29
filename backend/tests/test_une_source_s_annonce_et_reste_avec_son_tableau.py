"""Une source se lit comme une source, et ne quitte pas son tableau.

29/09/2026, business plan ÉCLORE : « données du projet » imprimé SEUL, en
italique de 8 points, sous chaque tableau et chaque chiffre clé — 73 pages
(`composants.note_source`). Un libellé nu ne dit pas qu'il est une source ; et
rien ne le liait à son tableau : un tableau qui finissait en bas de page
laissait sa source seule en haut de la suivante.

Tenu ici, sur le `.docx` écrit : la ligne se lit « Source : données du
projet », la dernière ligne du tableau la garde avec elle (`keep_with_next`),
et le chiffre clé dit aussi « Source : ». Les contre-épreuves : une source qui
s'annonce déjà n'est pas doublée, et le commentaire d'une FIGURE — ce qu'elle
apprend, pas d'où elle vient — n'est pas travesti en source.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from generation.chapitres.typographie import FINE_INSECABLE
from generation.rendu_word.depuis_json import rendre_etude

SOURCE = f"Source{FINE_INSECABLE}: "


def _rendre(tmp_path: Path, *blocs: dict[str, Any]) -> Any:
    etude = {
        "titre": "Business plan",
        "marque": {"nom": "ÉCLORE"},
        "chapitres": [{
            "numero": 1, "titre": "Le projet",
            "blocs": [
                {"type": "bandeau", "numero": 1, "titre": "Le projet", "accroche": ""},
                *blocs,
            ],
        }],
    }
    return Document(str(rendre_etude(etude, tmp_path / "sources.docx")))


def _tableau(source: str) -> dict[str, Any]:
    return {
        "type": "tableau",
        "entetes": ["Poste", "Montant"],
        "lignes": [["Loyer", "3 600 €"], ["Assurance", "480 €"]],
        "source": source,
    }


def _ce_qui_suit_le_tableau(document: Any, entete: str) -> tuple[Table, Paragraph]:
    elements = list(document.element.body.iterchildren())
    for rang, element in enumerate(elements):
        if element.tag.endswith("}tbl"):
            table = Table(element, document)
            if table.rows[0].cells[0].text == entete:
                return table, Paragraph(elements[rang + 1], document)
    raise AssertionError(f"tableau « {entete} » introuvable")


def test_la_source_d_un_tableau_s_annonce(tmp_path: Path) -> None:
    """Le défaut exact : « données du projet », seul, sous le tableau."""
    document = _rendre(tmp_path, _tableau("données du projet"))
    _table, source = _ce_qui_suit_le_tableau(document, "Poste")
    assert source.text == f"{SOURCE}données du projet"


def test_la_source_reste_avec_son_tableau(tmp_path: Path) -> None:
    """La dernière ligne du tableau garde la source sur sa page."""
    document = _rendre(tmp_path, _tableau("données du projet"))
    table, _source = _ce_qui_suit_le_tableau(document, "Poste")
    derniere = [p for cellule in table.rows[-1].cells for p in cellule.paragraphs]
    assert derniere and all(p.paragraph_format.keep_with_next for p in derniere)


def test_la_source_d_un_chiffre_cle_s_annonce(tmp_path: Path) -> None:
    document = _rendre(tmp_path, {
        "type": "kpi",
        "chiffres": [("5 000 €", "Budget de démarrage", "données du projet")],
    })
    cellules = [
        cellule.text for table in document.tables for ligne in table.rows
        for cellule in ligne.cells
    ]
    assert any(c.endswith(f"{SOURCE}données du projet") for c in cellules), cellules


def test_une_source_deja_annoncee_n_est_pas_doublee(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : pas de « Source : Source : Insee »."""
    document = _rendre(tmp_path, _tableau("Source : Insee, 2025"))
    _table, source = _ce_qui_suit_le_tableau(document, "Poste")
    assert source.text == "Source : Insee, 2025"


@pytest.mark.parametrize("texte", ["Selon la Fevad, 2025", "D'après Xerfi 2026"])
def test_une_attribution_ecrite_n_est_pas_prefixee(tmp_path: Path, texte: str) -> None:
    document = _rendre(tmp_path, _tableau(texte))
    _table, source = _ce_qui_suit_le_tableau(document, "Poste")
    assert source.text == texte


def test_un_tableau_sans_source_n_en_invente_pas(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : pas de « Source : » vide."""
    document = _rendre(tmp_path, _tableau(""))
    textes = [p.text for p in document.paragraphs]
    assert not any(t.startswith("Source") for t in textes), textes


def test_le_commentaire_d_une_figure_n_est_pas_une_source(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : sous une figure, le champ porte ce qu'elle APPREND.

    Le préfixer ferait d'une phrase d'analyse une attribution de source.
    """
    commentaire = "Le chiffre d'affaires double entre la première et la troisième année."
    document = _rendre(tmp_path, {
        "type": "graphique", "graphique": "barres", "titre": "Chiffre d'affaires",
        "source": commentaire,
        "donnees": {"etiquettes": ["2027", "2029"], "valeurs": [40.0, 80.0], "unite": " k€"},
    })
    textes = [p.text for p in document.paragraphs]
    assert commentaire in textes, textes
