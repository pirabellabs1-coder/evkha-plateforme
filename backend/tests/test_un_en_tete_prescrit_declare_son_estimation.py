"""Un tableau que le prompt prescrit doit déclarer ce que le contrôle exige.

## Le défaut mesuré

Corpus du 14/09/2026 : sur les études concurrentielles, la majorité des
« chiffres hors socle » restants sont les chiffres d'affaires projetés du
chapitre 6 — « CA de référence (2024, milieu de fourchette) : Botify Conseil |
260 000 € | 300 000 € | +15 % ». Le contrôle admet un montant ESTIMÉ qui montre
sa base ; il ne voyait ici aucune estimation déclarée.

Le prompt en était la cause : « Tableau : CA de référence, CA actuel,
évolution, commentaire court ». Il ordonnait des en-têtes que le contrôle
punissait — la règle 5 entre une consigne et un contrôle, déjà vécue avec
`[[UNDERSTAND]]`.

## Ce que ce fichier verrouille

- toute colonne de chiffre d'affaires prescrite par un prompt déclare son
  estimation, jugée par `_ESTIMATION` — l'expression qui FAIT FOI ;
- le milieu d'une fourchette et une projection sont des estimations ;
- contre-épreuve : un CA nu, sans rien qui le déclare, reste signalé.
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from generation.socle.referentiel import Fiabilite, Perimetre
from generation.socle.schema import DonneeSocle, Socle, Zone
from generation.verification.controles import _ESTIMATION, controler_chiffres_hors_socle
from generation.verification.lecture import lire_livrable

_COLONNE_DE_CA = re.compile(r"\bCA\b|chiffres?\s+d['’]affaires", re.IGNORECASE)


def _dossier_prompts() -> Path:
    from generation.chapitres.fichiers_prompts import rendre_prompt

    return Path(rendre_prompt.__code__.co_filename).parents[3] / "prompts"


def test_les_colonnes_de_ca_prescrites_declarent_leur_estimation() -> None:
    fautives: list[str] = []
    for fiche in sorted(_dossier_prompts().rglob("chapitre_*.md")):
        for consigne in re.findall(r"Tableau\s*:\s*([^.\n]+)", fiche.read_text(encoding="utf-8")):
            for colonne in consigne.split(","):
                if _COLONNE_DE_CA.search(colonne) and not _ESTIMATION.search(colonne):
                    fautives.append(f"{fiche.parent.name}/{fiche.name} : « {colonne.strip()} »")
    assert fautives == []


def _signales(tmp_path: Path, entetes: list[str], ligne: list[str]) -> set[str]:
    from docx import Document

    document = Document()
    document.add_paragraph("Projection des concurrents.")
    table = document.add_table(rows=2, cols=len(entetes))
    for colonne, (entete, valeur) in enumerate(zip(entetes, ligne, strict=True)):
        table.cell(0, colonne).text = entete
        table.cell(1, colonne).text = valeur
    chemin = tmp_path / "ec.docx"
    document.save(str(chemin))
    socle = Socle(
        secteur="conseil IA", zone=Zone(pays="France"), date_socle=date(2026, 9, 14),
        donnees=[DonneeSocle(
            id="marche_national_taille", libelle="Marché national", valeur=850,
            unite="MEUR", annee=2025, perimetre=Perimetre.NATIONAL,
            fiabilite=Fiabilite.OBSERVEE, source="Xerfi",
        )],
    )
    anomalies = controler_chiffres_hors_socle(lire_livrable(chemin), socle)
    return {a.detail.split(" »")[0].lstrip("« ") for a in anomalies}


def test_le_milieu_d_une_fourchette_est_une_estimation(tmp_path: Path) -> None:
    signales = _signales(
        tmp_path,
        ["Concurrent", "CA de référence (2024, milieu de fourchette)", "CA projeté 2026"],
        ["Botify Conseil", "260 000 €", "300 000 €"],
    )
    assert "260 000 €" not in signales
    assert "300 000 €" not in signales


def test_un_ca_que_rien_ne_declare_reste_signale(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : « CA de référence » seul n'est pas une estimation."""
    signales = _signales(
        tmp_path, ["Concurrent", "CA de référence", "CA actuel"],
        ["Botify Conseil", "260 000 €", "300 000 €"],
    )
    assert {"260 000 €", "300 000 €"} <= signales
