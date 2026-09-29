"""L'annexe des chiffres n'est pas un chapitre, et le suivi ne compte que les vendus.

Décision D9 du 29/09/2026 (`docs/diagnostic.md`, § 7 bis et § 8).

Business plan ÉCLORE : le sommaire du PDF livré comptait **22 entrées
numérotées** quand l'offre en annonce 21 — les chapitres 1 à 21 du plan, plus
« 22 — D'où viennent les chiffres de cette étude », l'annexe AJOUTÉE par le
rendu et numérotée comme un chapitre (`assemblage.py`, numéro du dernier
chapitre + 1). Le même +1 valait pour les quatre livrables. Et le suivi client
affichait « 22 chapitres sur 22 », parce qu'il comptait la Fiche projet.

Tenu ici :

- le document porte exactement ses chapitres numérotés, et l'annexe sous un
  bandeau « ANNEXE », sans numéro ;
- elle reste au sommaire, sans numéro, après les chapitres ;
- la lecture du document livré ne l'attribue à AUCUN chapitre ;
- le suivi affiche N = 21 pour un business plan (22, 9, 20 pour les autres).

Chacun échoue sur le code d'avant ; les contre-épreuves tiennent que l'annexe
n'a pas disparu.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from docx import Document

from generation.rendu_word.assemblage import assembler_etude
from generation.rendu_word.depuis_json import rendre_etude
from tests.test_lot3_livrable import _chapitre, _socle

pytestmark = pytest.mark.django_db

TITRE_ANNEXE = "D'où viennent les chiffres de cette étude"


def _business_plan(tmp_path: Path) -> tuple[dict[str, Any], Path]:
    etude, _ = assembler_etude(
        socle=_socle(),
        chapitres=[_chapitre(numero) for numero in range(1, 22)],
        titre="Business plan",
        marque={"nom": "ÉCLORE"},
    )
    return etude, rendre_etude(etude, tmp_path / "bp.docx")


def _premieres_cellules(chemin: Path) -> list[str]:
    return [
        table.rows[0].cells[0].text for table in Document(str(chemin)).tables
    ]


def test_le_document_compte_exactement_ses_vingt_et_un_chapitres(tmp_path: Path) -> None:
    """Le défaut exact : « CHAPITRE 22 » sur un business plan de 21 chapitres."""
    _etude, chemin = _business_plan(tmp_path)
    marqueurs = [
        texte.split("\n", 1)[0] for texte in _premieres_cellules(chemin)
        if texte.startswith("CHAPITRE ")
    ]
    assert marqueurs == [f"CHAPITRE {n:02d}" for n in range(1, 22)], marqueurs


def test_l_annexe_est_rendue_sous_un_bandeau_sans_numero(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : l'annexe n'a pas disparu, elle a changé de statut."""
    _etude, chemin = _business_plan(tmp_path)
    bandeaux = [t for t in _premieres_cellules(chemin) if t.startswith("ANNEXE\n")]
    assert len(bandeaux) == 1, _premieres_cellules(chemin)
    assert TITRE_ANNEXE.upper() in bandeaux[0]


def test_le_sommaire_liste_l_annexe_sans_numero(tmp_path: Path) -> None:
    _etude, chemin = _business_plan(tmp_path)
    sommaire = next(
        table for table in Document(str(chemin)).tables
        if table.rows[0].cells[0].text == "Chap."
    )
    lignes = [(ligne.cells[0].text, ligne.cells[1].text) for ligne in sommaire.rows[1:]]
    numeros = [numero for numero, _titre in lignes if numero]
    assert numeros == [f"{n:02d}" for n in range(1, 22)], numeros
    assert lignes[-1] == ("", f"Annexe — {TITRE_ANNEXE}")


def test_l_assemblage_range_l_annexe_hors_des_chapitres(tmp_path: Path) -> None:
    """Ce que la numérotation, la vérification et les comptes lisent."""
    etude, _chemin = _business_plan(tmp_path)
    assert [c["numero"] for c in etude["chapitres"]] == list(range(1, 22))
    assert [a["titre"] for a in etude["annexes"]] == [TITRE_ANNEXE]


def test_la_lecture_du_livrable_n_attribue_l_annexe_a_aucun_chapitre(tmp_path: Path) -> None:
    """Sans bandeau numéroté, l'annexe aurait été comptée au chapitre 21.

    Un défaut trouvé dans ses tableaux aurait alors envoyé réécrire « Sources et
    méthodologie », qui n'y est pour rien. Sur le code d'avant, elle était
    attribuée au chapitre 22 — un chapitre qui n'existe pas.
    """
    from generation.verification.lecture import lire_livrable

    _etude, chemin = _business_plan(tmp_path)
    lu = lire_livrable(chemin)
    chapitre_de_l_annexe = {
        chapitre for cellule, chapitre in zip(lu.cellules, lu.chapitre_de_la_cellule, strict=True)
        if cellule in ("Donnée", "Origine") or TITRE_ANNEXE.upper() in cellule
    }
    assert chapitre_de_l_annexe == {None}, chapitre_de_l_annexe
    assert 21 in lu.chapitre_de_la_cellule, "le chapitre 21 reste lu comme tel"
    assert 22 not in lu.chapitre_de_la_cellule


# ── Le suivi client ──────────────────────────────────────────────────────────


def _job_termine(type_livrable: str) -> Any:
    from catalog.models import Offer
    from customers.models import Customer
    from generation.models import ChapterGeneration, ChapterStatus
    from generation.services import bootstrap_generation_job
    from intake.models import IntakeStatus, IntakeSubmission
    from orders.models import Order

    offre = Offer.objects.create(
        name=type_livrable, slug=f"d9-{type_livrable}", deliverable_type=type_livrable,
    )
    client = Customer.objects.create(email=f"d9-{type_livrable}@example.com")
    commande = Order.objects.create(
        systeme_order_id=f"order-d9-{type_livrable}", customer=client, offer=offre,
    )
    soumission = IntakeSubmission.objects.create(
        order=commande, status=IntakeStatus.NORMALIZED,
        normalized_variables={"SECTEUR": "bien-être", "PAYS": "France",
                              "PROJET": "ÉCLORE", "ZONE": "Drôme"},
    )
    job = bootstrap_generation_job(soumission)
    ChapterGeneration.objects.filter(job=job).update(status=ChapterStatus.DONE)
    return job


@pytest.mark.parametrize(
    ("type_livrable", "annonces"),
    [
        ("business_plan", 21),
        ("market_study", 22),
        ("competitor_study", 9),
        ("business_strategy", 20),
    ],
)
def test_le_suivi_ne_compte_que_les_chapitres_annonces(
    type_livrable: str, annonces: int,
) -> None:
    """« 22 chapitres sur 22 » pour un business plan qui en annonce 21."""
    from organisations import suivi

    job = _job_termine(type_livrable)
    detail = next(e.detail for e in suivi.etapes(job) if e.cle == "chapitres")
    assert detail.startswith(f"{annonces} chapitres sur {annonces}"), detail


def test_le_suivi_garde_l_etat_de_la_production() -> None:
    """CONTRE-ÉPREUVE : la fiche projet reste du travail de production.

    Quand elle seule est écrite, l'étape « Rédaction des chapitres » est en
    cours — elle ne retombe pas en attente parce que le compte affiché est 0.
    """
    from generation.models import ChapterGeneration, ChapterStatus, JobStatus
    from organisations import suivi

    job = _job_termine("business_plan")
    job.status = JobStatus.RUNNING
    job.save(update_fields=["status"])
    ChapterGeneration.objects.filter(job=job, chapter_number__gte=1).update(
        status=ChapterStatus.PENDING,
    )
    etape = next(e for e in suivi.etapes(job) if e.cle == "chapitres")
    assert etape.etat == "en_cours"
    assert etape.detail.startswith("0 chapitre sur 21"), etape.detail
