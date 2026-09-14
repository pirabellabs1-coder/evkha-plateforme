"""Deux justifications que le lecteur voit et que le contrôle ne voyait pas.

## Les motifs, relevés sur le corpus du 14/09/2026 (études concurrentielles)

- « Évolution : EY France | 480 000 | 550 000 | +15 % » : la variation entre
  deux cellules de sa ligne. Deux nombres nus ne passaient pas pour un calcul
  POSÉ faute de mot de calcul dans la phrase — l'en-tête « Évolution » le dit.
- « Donnée : Prix moyen d'une baguette (fourchette observée 1,30 - 1,60 €,
  médiane retenue) | 1,45 EUR » : les bornes sont écrites dans le LIBELLÉ de la
  donnée du socle, que la référence ignorait.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from generation.socle.referentiel import Fiabilite, Perimetre
from generation.socle.schema import DonneeSocle, Socle, Zone
from generation.verification.controles import controler_chiffres_hors_socle
from generation.verification.lecture import lire_livrable


def _socle() -> Socle:
    return Socle(
        secteur="boulangerie", zone=Zone(pays="France"), date_socle=date(2026, 9, 14),
        donnees=[DonneeSocle(
            id="prix_baguette",
            libelle=(
                "Prix moyen d'une baguette tradition sur la zone "
                "(fourchette observée 1,30 - 1,60 €, médiane retenue)"
            ),
            valeur=1.45, unite="EUR", annee=2025, perimetre=Perimetre.NATIONAL,
            fiabilite=Fiabilite.ESTIMEE, source="EVKHA",
        )],
    )


def _signales(tmp_path: Path, tableau: list[list[str]], prose: str = "Synthèse.") -> set[str]:
    from docx import Document

    document = Document()
    document.add_paragraph(prose)
    table = document.add_table(rows=len(tableau), cols=len(tableau[0]))
    for rang, ligne in enumerate(tableau):
        for colonne, valeur in enumerate(ligne):
            table.cell(rang, colonne).text = valeur
    chemin = tmp_path / "ec.docx"
    document.save(str(chemin))
    anomalies = controler_chiffres_hors_socle(lire_livrable(chemin), _socle())
    return {a.detail.split(" »")[0].lstrip("« ") for a in anomalies}


def test_une_evolution_entre_deux_cellules_est_un_calcul(tmp_path: Path) -> None:
    signales = _signales(tmp_path, [
        ["Concurrent", "CA 2024", "CA 2026", "Évolution"],
        ["EY France", "480 000", "550 000", "+15 %"],
    ])
    assert "+15 %" not in signales and "15 %" not in signales


def test_une_evolution_fausse_reste_signalee(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : 480 000 → 550 000 ne fait pas +25 %."""
    signales = _signales(tmp_path, [
        ["Concurrent", "CA 2024", "CA 2026", "Évolution"],
        ["EY France", "480 000", "550 000", "+25 %"],
    ])
    assert any("25 %" in s for s in signales)


def test_sans_en_tete_de_resultat_deux_nombres_nus_ne_posent_rien(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : un taux quelconque ne devient pas un calcul par coïncidence."""
    signales = _signales(tmp_path, [
        ["Concurrent", "Clients", "Salariés", "Taux de fidélité"],
        ["EY France", "480 000", "550 000", "15 %"],
    ])
    assert any("15 %" in s for s in signales)


def test_les_bornes_ecrites_dans_le_libelle_du_socle_sont_du_socle(tmp_path: Path) -> None:
    signales = _signales(tmp_path, [
        ["Donnée", "Valeur", "Année"],
        ["Prix moyen d'une baguette (fourchette observée 1,30 - 1,60 €, médiane retenue)",
         "1,45 EUR", "2025"],
    ])
    assert "1,60 €" not in signales


def test_une_borne_absente_du_libelle_reste_signalee(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : 1,90 € n'est écrit nulle part dans le socle."""
    signales = _signales(tmp_path, [
        ["Donnée", "Valeur", "Année"],
        ["Prix moyen d'une baguette (fourchette observée 1,30 - 1,90 €)", "1,45 EUR", "2025"],
    ])
    assert "1,90 €" in signales
