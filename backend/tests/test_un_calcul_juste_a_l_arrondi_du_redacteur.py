"""Un calcul juste se reconnaît à l'arrondi qu'a choisi le rédacteur, et à son signe.

## Les motifs faux, relevés sur le corpus du 14/09/2026

- « SAM (1 824 M€) × 0,03 % ≈ 0,55 M€ » (étude de marché `ef567688`) :
  547 200 €, soit 0,51 % d'écart — au-delà de la tolérance fixe de 0,5 %. Or
  « 0,55 M€ » vaut à 5 000 € près : c'est l'arrondi du rédacteur, que le
  contrôle respectait déjà pour les pourcentages et pas pour les montants.
- « Dette résiduelle | 20 400 € | -6 600 € (20 400 - 27 000) » (`5c5e91b9`) et
  « Écart | 30 000 € | 15 000 € | -15 000 € » (`73dde3ab`) : seule la valeur
  ABSOLUE d'une différence était essayée.
- « Investissement total | 180 000 € | 100,0 % » sous « Part du total »
  (`b8da2640`) : la ligne du total vaut 100 % d'elle-même.

Les contre-épreuves gardent le contrôle armé : un résultat hors de l'arrondi,
une différence fausse, une part qui ne tombe pas juste restent signalés.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from generation.socle.referentiel import Fiabilite, Perimetre
from generation.socle.schema import DonneeSocle, Socle, Zone
from generation.verification.controles import controler_chiffres_hors_socle
from generation.verification.lecture import lire_livrable


def _socle() -> Socle:
    return Socle(
        secteur="bien-être", zone=Zone(pays="France"), date_socle=date(2026, 9, 14),
        donnees=[DonneeSocle(
            id="sam", libelle="Marché adressable", valeur=1824,
            unite="MEUR", annee=2025, perimetre=Perimetre.NATIONAL,
            fiabilite=Fiabilite.ESTIMEE, source="EVKHA",
        )],
    )


def _document(
    tmp_path: Path, paragraphes: list[str], tableau: list[list[str]] | None = None,
) -> Path:
    from docx import Document

    document = Document()
    for texte in paragraphes:
        document.add_paragraph(texte)
    if tableau:
        table = document.add_table(rows=len(tableau), cols=len(tableau[0]))
        for rang, ligne in enumerate(tableau):
            for colonne, valeur in enumerate(ligne):
                table.cell(rang, colonne).text = valeur
    chemin = tmp_path / "document.docx"
    document.save(str(chemin))
    return chemin


def _signales(chemin: Path, brief: tuple[float, ...] = ()) -> set[str]:
    anomalies = controler_chiffres_hors_socle(lire_livrable(chemin), _socle(), brief)
    return {a.detail.split(" »")[0].lstrip("« ") for a in anomalies}


def test_un_montant_arrondi_par_le_redacteur_est_un_calcul_juste(tmp_path: Path) -> None:
    chemin = _document(tmp_path, [
        "Année 1, montée en charge : SAM (1 824 M€) × 0,03 % ≈ 0,55 M€, taux prudent.",
    ])
    assert "0,55 M€" not in _signales(chemin)


def test_un_montant_hors_de_son_arrondi_reste_signale(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : 547 200 € ne s'écrit pas 0,60 M€."""
    chemin = _document(tmp_path, [
        "Année 1, montée en charge : SAM (1 824 M€) × 0,03 % ≈ 0,60 M€, taux prudent.",
    ])
    assert "0,60 M€" in _signales(chemin)


def test_une_difference_negative_posee_est_un_calcul(tmp_path: Path) -> None:
    chemin = _document(tmp_path, ["Synthèse de la dette."], [
        ["Poste", "Exercice 2", "Évolution vs exercice 1"],
        ["Dette résiduelle", "20 400 €", "-6 600 € (20 400 - 27 000)"],
    ])
    assert "-6 600 €" not in _signales(chemin, (20400.0, 27000.0))


def test_une_difference_negative_fausse_reste_signalee(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : 20 400 - 27 000 ne fait pas -8 600."""
    chemin = _document(tmp_path, ["Synthèse de la dette."], [
        ["Poste", "Exercice 2", "Évolution vs exercice 1"],
        ["Dette résiduelle", "20 400 €", "-8 600 € (20 400 - 27 000)"],
    ])
    assert "-8 600 €" in _signales(chemin, (20400.0, 27000.0))


def test_la_ligne_du_total_vaut_cent_pour_cent(tmp_path: Path) -> None:
    chemin = _document(tmp_path, ["Plan de financement."], [
        ["Poste", "Montant", "Part du total"],
        ["Apport", "45 000 €", "25,0 %"],
        ["Investissement total", "180 000 €", "100,0 %"],
    ])
    signales = _signales(chemin, (180000.0, 45000.0))
    assert "100,0 %" not in signales
    assert "25,0 %" not in signales


def test_une_part_qui_ne_tombe_pas_juste_reste_signalee(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : 45 000 / 180 000 ne fait pas 30 %."""
    chemin = _document(tmp_path, ["Plan de financement."], [
        ["Poste", "Montant", "Part du total"],
        ["Apport", "45 000 €", "30,0 %"],
    ])
    assert "30,0 %" in _signales(chemin, (180000.0, 45000.0))


def test_seule_la_ligne_du_total_vaut_cent_pour_cent(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : 45 000 € au brief ne fait pas de l'apport 100 % du plan."""
    chemin = _document(tmp_path, ["Plan de financement."], [
        ["Poste", "Montant", "Part du total"],
        ["Apport", "45 000 €", "100,0 %"],
    ])
    assert "100,0 %" in _signales(chemin, (180000.0, 45000.0))


def test_un_nombre_ecrit_sans_decimale_n_ouvre_pas_une_tolerance_geante(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : 2 600 000 + 12 ne s'arrondit pas en « 3 M€ » de ventes."""
    chemin = _document(tmp_path, [
        "Sur 2 600 000 habitants, 12 % achètent, soit 3 M€ de ventes.",
    ])
    assert "3 M€" in _signales(chemin)
