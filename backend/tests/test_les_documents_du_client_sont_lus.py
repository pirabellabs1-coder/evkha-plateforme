"""Les documents déposés par le client sont lus par la génération.

Dossier `f7f2fad9` (stratégie, 08/09/2026). La cliente avait joint son
prévisionnel et une étude de marché régionale. Aucun livrable ne lisait les
pièces jointes : le document a posé un marché national « estimé à 900 millions
d'euros », sans source, et pris le chiffre d'affaires actuel pour un objectif.

Ces tests échouent sur le code d'avant (règle 6) : `generation.documents_client`
n'existait pas, et ni le socle, ni sa vérification, ni les chapitres ne
recevaient autre chose que le brief et la recherche web. Les contre-épreuves
vérifient qu'un dossier SANS document garde exactement ses prompts d'avant, et
qu'un document illisible n'est jamais présenté au modèle comme lu.
"""
from __future__ import annotations

import zipfile
from datetime import date
from io import BytesIO
from typing import Any

import pytest
from django.core.files.base import ContentFile

from generation import documents_client as dc
from generation.models import DocumentClientLu, StatutLecture

# ── Fabriques de fichiers ────────────────────────────────────────────────────

_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _docx(paragraphes: list[str], tableau: list[list[str]] | None = None) -> bytes:
    corps = "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in paragraphes)
    if tableau:
        corps += "<w:tbl>" + "".join(
            "<w:tr>" + "".join(
                f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in ligne
            ) + "</w:tr>"
            for ligne in tableau
        ) + "</w:tbl>"
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        z.writestr("word/document.xml", f"<w:document {_W}><w:body>{corps}</w:body></w:document>")
    return tampon.getvalue()


_S = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
_R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


def _xlsx() -> bytes:
    """Deux feuilles, chaînes partagées, nombres, colonne sautée, texte en ligne."""
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        z.writestr("xl/workbook.xml", (
            f"<workbook {_S} {_R}><sheets>"
            '<sheet name="Prévisionnel" sheetId="1" r:id="rId1"/>'
            '<sheet name="Formules" sheetId="2" r:id="rId2"/>'
            "</sheets></workbook>"
        ))
        z.writestr("xl/_rels/workbook.xml.rels", (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="x" Target="worksheets/sheet1.xml"/>'
            '<Relationship Id="rId2" Type="x" Target="/xl/worksheets/sheet2.xml"/>'
            "</Relationships>"
        ))
        z.writestr("xl/sharedStrings.xml", (
            f"<sst {_S}><si><t>Chiffre d'affaires</t></si><si><t>Objectif 2028</t></si>"
            "<si><r><t>Abonnés </t></r><r><t>Essentiel</t></r></si></sst>"
        ))
        z.writestr("xl/worksheets/sheet1.xml", (
            f"<worksheet {_S}><sheetData>"
            '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1"><v>3000</v></c>'
            '<c r="D1"><v>0.20000000000000001</v></c></row>'
            '<row r="2"><c r="A2" t="s"><v>1</v></c><c r="B2"><v>45000</v></c></row>'
            "</sheetData></worksheet>"
        ))
        z.writestr("xl/worksheets/sheet2.xml", (
            f"<worksheet {_S}><sheetData>"
            '<row r="1"><c r="A1" t="s"><v>2</v></c>'
            '<c r="B1" t="inlineStr"><is><t>9</t></is></c></row>'
            "</sheetData></worksheet>"
        ))
    return tampon.getvalue()


def _pptx() -> bytes:
    a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        for numero, texte in ((10, "Dixième"), (2, "Deuxième"), (1, "Première")):
            z.writestr(
                f"ppt/slides/slide{numero}.xml",
                f"<p:sld {a} xmlns:p=\"p\"><a:p><a:r><a:t>{texte}</a:t></a:r></a:p></p:sld>",
            )
    return tampon.getvalue()


def _pdf(lignes: list[str]) -> bytes:
    from reportlab.pdfgen import canvas

    tampon = BytesIO()
    toile = canvas.Canvas(tampon)
    for rang, ligne in enumerate(lignes):
        toile.drawString(72, 760 - 14 * rang, ligne)
    toile.showPage()
    toile.save()
    return tampon.getvalue()


def _pdf_sans_texte() -> bytes:
    from reportlab.pdfgen import canvas

    tampon = BytesIO()
    toile = canvas.Canvas(tampon)
    toile.rect(50, 50, 200, 200, fill=1)  # une « image » de page, aucun texte
    toile.showPage()
    toile.save()
    return tampon.getvalue()


# ── Extraction, format par format ────────────────────────────────────────────


def test_un_document_word_est_lu_avec_ses_tableaux() -> None:
    lu = dc.extraire("prev.docx", _docx(
        ["Objectif : 250 abonnés à 24 mois."],
        [["Formule", "Prix", "Abonnés"], ["Essentiel", "12 €", "9"]],
    ))
    assert lu.statut == StatutLecture.LU
    assert "250 abonnés à 24 mois" in lu.texte
    assert "Essentiel | 12 € | 9" in lu.texte, "une ligne de tableau reste une ligne"


def test_un_classeur_excel_est_lu_feuille_par_feuille() -> None:
    lu = dc.extraire("previsionnel.xlsx", _xlsx())
    assert lu.statut == StatutLecture.LU, lu.motif
    assert "### Feuille « Prévisionnel »" in lu.texte
    # Colonne C vide conservée : sans elle, la valeur de D glisserait sous C.
    # Sans feuille de styles, un nombre reste un nombre, écrit à la française.
    assert "Chiffre d'affaires | 3000 | | 0,2" in lu.texte
    assert "Objectif 2028 | 45000" in lu.texte
    # Chaîne partagée en plusieurs morceaux, texte en ligne, cible absolue.
    assert "Abonnés Essentiel | 9" in lu.texte


def test_une_presentation_se_lit_dans_l_ordre_des_diapositives() -> None:
    lu = dc.extraire("deck.pptx", _pptx())
    assert lu.statut == StatutLecture.LU
    assert lu.texte.index("Première") < lu.texte.index("Deuxième") < lu.texte.index("Dixième"), (
        "slide10 trié comme texte passerait avant slide2"
    )


def test_un_pdf_texte_est_lu() -> None:
    lu = dc.extraire("etude.pdf", _pdf([
        "Etude regionale du marche de l'assistance informatique, CCI 2025.",
        "Le marche regional est estime a 14,2 millions d'euros en 2024.",
        "Il progresse de 4 % par an selon la meme source et concerne 38 000 foyers.",
    ]))
    assert lu.statut == StatutLecture.LU, lu.motif
    assert "14,2 millions" in lu.texte


@pytest.mark.parametrize(
    ("nom", "contenu", "statut", "dans_le_motif"),
    [
        ("scan.pdf", "pdf_sans_texte", StatutLecture.ILLISIBLE, "scann"),
        ("ancien.doc", b"\xd0\xcf\x11\xe0", StatutLecture.NON_LU, ".docx"),
        ("ancien.xls", b"\xd0\xcf\x11\xe0", StatutLecture.NON_LU, ".xlsx"),
        ("photo.png", b"\x89PNG", StatutLecture.NON_LU, "image"),
        ("casse.docx", b"PK\x03\x04pas une archive", StatutLecture.ILLISIBLE, "endommag"),
    ],
)
def test_ce_qui_n_est_pas_lu_le_dit(
    nom: str, contenu: Any, statut: str, dans_le_motif: str
) -> None:
    """Règle 1 : un document écarté porte sa raison, il ne passe pas pour lu."""
    octets = _pdf_sans_texte() if contenu == "pdf_sans_texte" else contenu
    lu = dc.extraire(nom, octets)
    assert lu.statut == statut
    assert lu.texte == ""
    assert dans_le_motif in lu.motif


def test_une_partie_avec_declaration_d_entites_est_refusee() -> None:
    """Une partie Office n'a jamais de DTD : une DTD, c'est une bombe possible."""
    bombe = (
        '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa">'
        '<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]>'
        f"<w:document {_W}><w:body><w:p><w:r><w:t>&b;</w:t></w:r></w:p></w:body></w:document>"
    )
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        z.writestr("word/document.xml", bombe)
    lu = dc.extraire("piege.docx", tampon.getvalue())
    assert lu.statut == StatutLecture.ILLISIBLE
    assert "refusée" in lu.motif


# ── Répartition du volume ────────────────────────────────────────────────────


def test_un_long_document_ne_prive_pas_les_autres_de_lecture() -> None:
    parts = dc.repartir([100_000, 3_000, 500], 10_000)
    assert parts[1] == 3_000 and parts[2] == 500, "les courts sont servis en entier"
    assert parts[0] == 10_000 - 3_500, "le long prend ce qui reste"
    assert sum(dc.repartir([50_000, 50_000], 10_000)) == 10_000


# ── Le dossier ───────────────────────────────────────────────────────────────


@pytest.fixture(autouse=True)
def media_isole(tmp_path: Any, settings: Any) -> None:
    settings.MEDIA_ROOT = str(tmp_path)


@pytest.fixture
def job_avec_organisation(db: Any) -> Any:
    from catalog.models import DeliverableType, Offer
    from customers.models import Customer
    from generation.services import bootstrap_generation_job
    from intake.models import IntakeStatus, IntakeSubmission
    from orders.models import Order
    from organisations import services

    contact = Customer.objects.create(email="documents@exemple.fr")
    organisation = services.creer_organisation(raison_sociale="Atelier Test", contact=contact)
    offre = Offer.objects.create(
        name="Stratégie", slug="strat-docs",
        deliverable_type=DeliverableType.BUSINESS_STRATEGY,
    )
    commande = Order.objects.create(
        systeme_order_id="cmd-docs", customer=contact, offer=offre,
        organisation=organisation,
    )
    soumission = IntakeSubmission.objects.create(
        order=commande, status=IntakeStatus.NORMALIZED,
        normalized_variables={"SECTEUR": "assistance informatique", "PAYS": "France"},
    )
    return bootstrap_generation_job(soumission)


def _deposer(job: Any, nom: str, contenu: bytes, categorie: str = "document") -> Any:
    from organisations.models import PieceJointe

    organisation = job.order.organisation
    piece = PieceJointe(
        organisation=organisation, categorie=categorie, nom_original=nom,
        taille_octets=len(contenu),
    )
    piece.fichier.save(nom, ContentFile(contenu), save=False)
    piece.save()
    return piece


def test_la_bibliotheque_est_lue_et_tracee_sur_le_dossier(job_avec_organisation: Any) -> None:
    from monitoring.models import OperationalIncident

    job = job_avec_organisation
    _deposer(job, "previsionnel.xlsx", _xlsx())
    _deposer(job, "vieux.doc", b"\xd0\xcf\x11\xe0")
    _deposer(job, "logo.png", b"\x89PNG", categorie="logo")

    lignes = dc.lire_les_documents(job)

    assert [(ligne.nom, ligne.statut) for ligne in lignes] == [
        ("previsionnel.xlsx", StatutLecture.LU),
        ("vieux.doc", StatutLecture.NON_LU),
    ], "le logo n'est pas un document de travail"
    assert "45000" in dc.bloc_documents(job)
    # Le document écarté n'est PAS nommé au modèle : il écrirait au client
    # qu'un fichier n'a pas pu être lu.
    assert "vieux.doc" not in dc.bloc_documents(job)
    incident = OperationalIncident.objects.get(job=job)
    assert "1 non lu(s)" in incident.title
    assert incident.details["documents"][0]["nom"] == "vieux.doc"


def test_une_relance_relit_la_meme_matiere(job_avec_organisation: Any) -> None:
    job = job_avec_organisation
    _deposer(job, "a.docx", _docx(["Premier document."]))
    dc.lire_les_documents(job)
    _deposer(job, "b.docx", _docx(["Ajouté entre deux passages."]))

    dc.lire_les_documents(job)

    assert job.documents_client.count() == 1
    assert "Ajouté entre deux passages" not in dc.bloc_documents(job)


def test_la_coupure_se_dit(job_avec_organisation: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(dc, "LIMITE_TOTALE", 60)
    job = job_avec_organisation
    _deposer(job, "long.docx", _docx(["x" * 50 + " fin"] * 10))

    (ligne,) = dc.lire_les_documents(job)

    assert ligne.statut == StatutLecture.TRONQUE
    assert ligne.caracteres_retenus <= 60 < ligne.caracteres_extraits
    # La coupure se dit dans la trace — jamais au modèle, qui l'écrirait.
    assert "transmis sur" in ligne.motif
    assert "Extrait" not in dc.bloc_documents(job)


def test_sans_organisation_rien_n_est_lu(db: Any) -> None:
    """Contre-épreuve : le flux Systeme.io n'a pas de bibliothèque."""
    from catalog.models import DeliverableType, Offer
    from customers.models import Customer
    from generation.services import bootstrap_generation_job
    from intake.models import IntakeStatus, IntakeSubmission
    from orders.models import Order

    offre = Offer.objects.create(
        name="EM", slug="em-docs", deliverable_type=DeliverableType.MARKET_STUDY
    )
    commande = Order.objects.create(
        systeme_order_id="cmd-sans-org",
        customer=Customer.objects.create(email="s@e.fr"),
        offer=offre,
    )
    soumission = IntakeSubmission.objects.create(
        order=commande, status=IntakeStatus.NORMALIZED, normalized_variables={"SECTEUR": "x"},
    )
    job = bootstrap_generation_job(soumission)
    assert dc.lire_les_documents(job) == []
    assert dc.bloc_documents(job) == ""


def test_supprimer_la_piece_efface_le_texte_mais_garde_la_trace(job_avec_organisation: Any) -> None:
    """La rétention de douze mois vaut pour le CONTENU, pas seulement le fichier."""
    job = job_avec_organisation
    piece = _deposer(job, "bilan.docx", _docx(["Résultat net : 12 400 €."]))
    dc.lire_les_documents(job)
    # Un dossier TERMINÉ : celui qui travaille garde son texte jusqu'à la fin
    # (voir `test_retirer_un_document_pendant_la_generation…`).
    job.status = "done"
    job.save(update_fields=["status"])

    piece.delete()

    ligne = DocumentClientLu.objects.get(job=job)
    assert ligne.texte == ""
    assert ligne.texte_efface_le is not None
    assert ligne.nom == "bilan.docx", "la ligne dit encore ce qui a été lu"
    assert dc.bloc_documents(job) == ""


# ── Les documents arrivent là où les chiffres se décident ────────────────────


def test_le_socle_recoit_les_documents_et_la_consigne() -> None:
    from generation.socle.prompt import CONSIGNE_DOCUMENTS_SOCLE, construire_prompt_socle

    avec = construire_prompt_socle(
        deliverable_type="business_strategy", variables={"SECTEUR": "x"},
        documents_client="DOCUMENTS_DU_CLIENT — étude régionale : 14,2 M€ (CCI, 2025)",
    )
    sans = construire_prompt_socle(deliverable_type="business_strategy", variables={"SECTEUR": "x"})

    assert "14,2 M€ (CCI, 2025)" in avec
    assert CONSIGNE_DOCUMENTS_SOCLE in avec
    assert "« données du projet »" in avec
    assert "c'est ELLE que tu retiens" in avec, "le brief l'emporte sur un document"
    assert "plus récente" not in CONSIGNE_DOCUMENTS_SOCLE
    # Contre-épreuve : sans document, le prompt d'avant, octet pour octet
    # (hormis la date du jour, identique dans les deux appels).
    assert CONSIGNE_DOCUMENTS_SOCLE not in sans
    assert avec.replace(
        "\n\nDOCUMENTS_DU_CLIENT — étude régionale : 14,2 M€ (CCI, 2025)\n\n"
        + CONSIGNE_DOCUMENTS_SOCLE, ""
    ) == sans


class _ClientEspion:
    def __init__(self) -> None:
        self.prompts: list[str] = []

    def complete_structured(self, **kwargs: Any) -> Any:
        self.prompts.append(kwargs["prompt"])

        class _R:
            payload = {"verdicts": [{"identifiant": "taille_marche", "statut": "confirmee"}]}

        return _R()


def _socle_observe() -> Any:
    from generation.socle.referentiel import Fiabilite, Perimetre
    from generation.socle.schema import DonneeSocle, Socle, Zone

    return Socle(
        secteur="assistance informatique", zone=Zone(pays="France", region="Bretagne"),
        date_socle=date(2026, 9, 1),
        donnees=[DonneeSocle(
            id="taille_marche", libelle="Marché régional", valeur=14.2, unite="MEUR",
            annee=2024, perimetre=Perimetre.REGIONAL, fiabilite=Fiabilite.OBSERVEE,
            source="CCI Bretagne, 2025",
        )],
    )


def test_la_verification_confronte_aussi_aux_documents() -> None:
    """Sans web, un chiffre publié qu'un document reproduit n'est plus déclassé d'office."""
    from generation.socle.referentiel import Fiabilite
    from generation.socle.verification import verifier_le_socle

    socle = _socle_observe()
    espion = _ClientEspion()
    rapport = verifier_le_socle(
        socle, client=espion, brief_recherche="",
        documents_client="DOCUMENTS_DU_CLIENT — CCI Bretagne 2025 : 14,2 M€",
    )

    assert len(espion.prompts) == 1, "l'ancien code déclassait sans même regarder"
    assert "CCI Bretagne 2025 : 14,2 M€" in espion.prompts[0]
    assert "SEULEMENT s'il l'attribue" in espion.prompts[0]
    assert rapport.confirmees == ["taille_marche"]
    assert socle.donnees[0].fiabilite == Fiabilite.OBSERVEE


def test_sans_web_ni_document_on_declasse_toujours() -> None:
    """Contre-épreuve de la règle 1 : rien à comparer, rien de confirmé."""
    from generation.socle.referentiel import Fiabilite
    from generation.socle.verification import verifier_le_socle

    socle = _socle_observe()
    espion = _ClientEspion()
    verifier_le_socle(socle, client=espion, brief_recherche="", documents_client="")
    assert espion.prompts == []
    assert socle.donnees[0].fiabilite == Fiabilite.ESTIMEE


def test_le_prompt_du_chapitre_porte_les_documents_dans_sa_partie_en_cache(
    job_avec_organisation: Any,
) -> None:
    from generation.chapitres.configuration import type_document
    from generation.chapitres.runner import CONSIGNE_DOCUMENTS_CHAPITRE, construire_prompt_chapitre

    job = job_avec_organisation
    _deposer(job, "repartition.docx", _docx(
        ["Répartition des abonnés"], [["Essentiel", "12 €", "9"], ["Confort", "19 €", "4"]]
    ))
    dc.lire_les_documents(job)
    chapitres = list(job.chapters.order_by("chapter_number")[:2])
    variables = {"SECTEUR": "assistance informatique", "PAYS": "France"}
    document = type_document(str(job.deliverable_type))

    prompts = [
        construire_prompt_chapitre(
            c, socle=_socle_observe(), variables=variables, document=document
        )[0]
        for c in chapitres
    ]

    for prompt in prompts:
        assert "Confort | 19 € | 4" in prompt.par_job
        assert CONSIGNE_DOCUMENTS_CHAPITRE in prompt.par_job
        assert "Confort | 19 € | 4" not in prompt.par_chapitre
    assert prompts[0].par_job == prompts[1].par_job, "un octet qui varie casse le cache"


def test_les_chiffres_des_documents_ne_sont_pas_signales_hors_socle(
    job_avec_organisation: Any,
) -> None:
    from generation.verification.services import chiffres_du_brief

    job = job_avec_organisation
    avant = chiffres_du_brief(job)
    _deposer(job, "prev.docx", _docx(["Chiffre d'affaires visé : 45 000 €."]))
    dc.lire_les_documents(job)

    assert 45000.0 in chiffres_du_brief(job)
    assert 45000.0 not in avant


def test_le_runner_lit_les_documents_avant_le_socle(
    job_avec_organisation: Any, monkeypatch: pytest.MonkeyPatch, settings: Any,
) -> None:
    """Le branchement réel : lus au lancement, transmis au socle."""
    from generation import runner as moteur
    from generation.chapitres import services as cycle_de_vie

    settings.EVKHA_SOCLE_ENABLED = True
    job = job_avec_organisation
    _deposer(job, "etude.docx", _docx(["Marché régional : 14,2 M€ selon la CCI, 2025."]))
    appels: list[dict[str, Any]] = []

    def _socle(job_recu: Any, **kwargs: Any) -> Any:
        appels.append(kwargs)
        return object()

    def _chapitre(job_recu: Any, numero: int, **kwargs: Any) -> Any:
        chapitre = job_recu.chapters.get(chapter_number=numero)
        chapitre.status = "done"
        chapitre.payload = {"chapitre": numero}
        chapitre.content = "# chapitre"
        chapitre.save(update_fields=["status", "payload", "content", "updated_at"])
        return chapitre

    monkeypatch.setattr(moteur, "debiter_pour_job", lambda job: (True, ""))
    monkeypatch.setattr(moteur, "etablir_socle", _socle)
    monkeypatch.setattr(moteur, "socle_verrouille", lambda job: object())
    monkeypatch.setattr(cycle_de_vie, "produire_chapitre", _chapitre)

    moteur.run_generation_job(job)

    assert "14,2 M€ selon la CCI" in appels[0]["documents_client"]
    assert job.documents_client.count() == 1


def test_la_console_liste_les_documents_sans_leur_texte(
    job_avec_organisation: Any, client_admin: Any,
) -> None:
    job = job_avec_organisation
    _deposer(job, "bilan.docx", _docx(["Résultat net confidentiel : 12 400 €."]))
    _deposer(job, "scan.pdf", _pdf_sans_texte())
    dc.lire_les_documents(job)

    reponse = client_admin.get(f"/api/dashboard/jobs/{job.id}/")

    assert reponse.status_code == 200
    documents = reponse.json()["documents_client"]
    assert [d["statut"] for d in documents] == ["lu", "illisible"]
    assert "scann" in documents[1]["motif"]
    assert "12 400" not in reponse.content.decode(), "le bilan du client ne sort pas"


# ── Bombes de décompression (audit du 11/09/2026) ────────────────────────────


def test_le_plafond_porte_sur_les_octets_decompresses_pas_sur_la_declaration(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Une partie qui dépasse le plafond est lue JUSQU'AU plafond, et le dit.

    L'ancien code comparait `ZipInfo.file_size`, écrit par l'expéditeur, puis
    décompressait tout en mémoire. La première correction refusait alors le
    document entier — une seule grosse partie rendait un prévisionnel
    illisible. Ce qui a été lu reste lu ; la coupure se déclare.
    """
    monkeypatch.setattr(dc, "MAX_OCTETS_ARBRE", 5_000)
    paragraphes = [f"Paragraphe {i} du rapport annuel." for i in range(2000)]
    lu = dc.extraire("gros.docx", _docx(paragraphes))
    assert lu.statut == StatutLecture.TRONQUE
    assert "lue en partie" in lu.motif
    assert "Paragraphe 0 du rapport annuel." in lu.texte
    assert "Paragraphe 1999" not in lu.texte


def test_une_archive_aux_milliers_d_entrees_est_refusee(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(dc, "MAX_ENTREES_ARCHIVE", 2)
    lu = dc.extraire("deck.pptx", _pptx())
    assert lu.statut == StatutLecture.ILLISIBLE
    assert "entrées" in lu.motif


def test_le_budget_vaut_pour_toute_l_archive(monkeypatch: pytest.MonkeyPatch) -> None:
    """Chaque partie sous le plafond, mais leur somme au-dessus : la fin est coupée."""
    classeur = _xlsx()
    with zipfile.ZipFile(BytesIO(classeur)) as z:
        taille = {i.filename: i.file_size for i in z.infolist()}
    # Tout, sauf la seconde moitié de la dernière feuille lue.
    budget = sum(taille.values()) - taille["xl/worksheets/sheet2.xml"] // 2
    monkeypatch.setattr(dc, "MAX_OCTETS_ARCHIVE", budget)
    lu = dc.extraire("previsionnel.xlsx", classeur)
    assert lu.statut == StatutLecture.TRONQUE
    assert "xl/worksheets/sheet2.xml lue en partie" in lu.motif
    assert "Objectif 2028 | 45000" in lu.texte, "la première feuille reste lue"


# ── Un webhook ne forge ni la gratuité ni le silence ─────────────────────────


def test_un_marqueur_de_reprise_venu_d_un_webhook_ne_vaut_rien(
    job_avec_organisation: Any, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Identifiant ET `raw_payload` d'une commande Systeme.io viennent du webhook.

    Les deux marques y étaient forgeables : exemption de débit et suppression
    de l'envoi au client. La preuve exigée désormais — un dossier d'origine
    réel du même client — n'est connue que du serveur.
    """
    import uuid

    from delivery import tasks as livraison
    from generation import tasks
    from organisations.liaison import est_une_reprise_a_nos_frais

    job = job_avec_organisation
    faux = str(uuid.uuid4())
    commande = job.order
    commande.systeme_order_id = f"reprise-{faux[:8]}-deadbeef"
    commande.raw_payload = {"reprise_de": faux, "sans_envoi": True}
    commande.save(update_fields=["systeme_order_id", "raw_payload"])

    assert not est_une_reprise_a_nos_frais(job)

    envois: list[str] = []
    monkeypatch.setattr(livraison.deliver_job_task, "delay", envois.append)
    tasks._livrer(job)
    assert envois == [str(job.id)], "le client reçoit son document"


def _bombe_qui_ment(decompresse: int, declare: int) -> bytes:
    """Un `.docx` dont la partie déclare `declare` octets et en produit `decompresse`."""
    import struct

    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b" " * decompresse)
    octets = bytearray(tampon.getvalue())
    for signature, decalage in ((b"PK\x03\x04", 22), (b"PK\x01\x02", 24)):
        position = octets.find(signature)
        struct.pack_into("<I", octets, position + decalage, declare)
    return bytes(octets)


def test_une_archive_qui_ment_sur_sa_taille_ne_gonfle_pas_la_memoire(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """LE test de l'audit : déclaré 100 octets, 60 Mo en réalité.

    L'ancien code croyait la déclaration, puis `archive.read()` décompressait
    tout d'un bloc : le pic mémoire suivait la taille RÉELLE (616 Mo mesurés
    par l'audit sur 300 Mo). Lue par morceaux bornés, la même archive ne
    dépasse jamais le plafond, quoi qu'elle déclare.
    """
    import tracemalloc

    monkeypatch.setattr(dc, "MAX_OCTETS_PARTIE", 1024 * 1024)
    bombe = _bombe_qui_ment(decompresse=60 * 1024 * 1024, declare=100)

    tracemalloc.start()
    lu = dc.extraire("piege.docx", bombe)
    _, pic = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    assert lu.statut == StatutLecture.ILLISIBLE
    assert pic < 20 * 1024 * 1024, f"pic mémoire de {pic // (1024 * 1024)} Mo"


# ── Relecture du 11/09/2026 : ce que les lecteurs perdaient ou fabriquaient ──


def test_une_tabulation_ne_colle_pas_l_annee_au_montant() -> None:
    """B1 : « 2025⇥45 000 € » devenait « 202545 000 € », un faux chiffre du client."""
    import docx
    from docx.shared import Inches

    from core.numbers import amounts_in

    document = docx.Document()
    paragraphe = document.add_paragraph()
    # Une tabulation de PARAGRAPHE (position) n'est pas un caractère.
    paragraphe.paragraph_format.tab_stops.add_tab_stop(Inches(3))
    run = paragraphe.add_run("Chiffre d'affaires 2025")
    run.add_tab()
    run.add_text("45 000 €")
    saut = document.add_paragraph().add_run("Objectif 2027")
    saut.add_break()
    saut.add_text("120 000 €")
    tampon = BytesIO()
    document.save(tampon)

    lu = dc.extraire("prev.docx", tampon.getvalue())

    assert "Chiffre d'affaires 2025 | 45 000 €" in lu.texte
    assert "Objectif 2027\n120 000 €" in lu.texte
    assert 202545000.0 not in amounts_in(lu.texte)
    assert 45000.0 in amounts_in(lu.texte)


def test_un_saut_de_ligne_de_diapositive_separe_les_nombres() -> None:
    a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        z.writestr(
            "ppt/slides/slide1.xml",
            f'<p:sld {a} xmlns:p="p"><a:p><a:r><a:t>Clients 2025</a:t></a:r>'
            "<a:br/><a:r><a:t>1 200</a:t></a:r></a:p></p:sld>",
        )
    assert "Clients 2025\n1 200" in dc.extraire("d.pptx", tampon.getvalue()).texte


def _xlsx_mis_en_forme() -> bytes:
    """Styles réels : pourcentage, date intégrée, date personnalisée ; cible `../`."""
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        z.writestr("xl/workbook.xml", (
            f"<workbook {_S} {_R}><sheets>"
            '<sheet name="Mensuel" sheetId="1" r:id="rId1"/>'
            '<sheet name="Fantôme" sheetId="2" r:id="rId2"/>'
            "</sheets></workbook>"
        ))
        z.writestr("xl/_rels/workbook.xml.rels", (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="x" Target="../xl/worksheets/sheet1.xml"/>'
            '<Relationship Id="rId2" Type="x" Target="worksheets/absente.xml"/>'
            "</Relationships>"
        ))
        z.writestr("xl/styles.xml", (
            f"<styleSheet {_S}><numFmts>"
            '<numFmt numFmtId="164" formatCode="dd/mm/yyyy"/>'
            '<numFmt numFmtId="165" formatCode="0.0&quot; mois&quot;"/>'
            "</numFmts><cellXfs>"
            '<xf numFmtId="0"/><xf numFmtId="9"/><xf numFmtId="14"/>'
            '<xf numFmtId="164"/><xf numFmtId="165"/>'
            "</cellXfs></styleSheet>"
        ))
        z.writestr("xl/worksheets/sheet1.xml", (
            f"<worksheet {_S}><sheetData>"
            '<row r="1"><c r="A1" t="inlineStr"><is><t>Marge</t></is></c>'
            '<c r="B1" s="1"><v>0.35</v></c><c r="C1" s="2"><v>46023</v></c>'
            '<c r="D1" s="3"><v>46054</v></c><c r="E1" s="4"><v>3.5</v></c>'
            '<c r="F1"><f>SUM(B1)</f></c><c r="EA1"><v>7</v></c></row>'
            "</sheetData></worksheet>"
        ))
    return tampon.getvalue()


def test_un_tableur_rend_ses_dates_et_pourcentages_et_declare_ses_pertes() -> None:
    lu = dc.extraire("prev.xlsx", _xlsx_mis_en_forme())

    assert "Marge | 35 % | 01/01/2026 | 01/02/2026 | 3,5" in lu.texte, lu.texte
    # Toute perte fait « lu en partie », avec sa raison.
    assert lu.statut == StatutLecture.TRONQUE
    assert "« Fantôme » introuvable" in lu.motif
    assert "1 cellule(s) calculée(s) sans valeur" in lu.motif
    assert "colonnes au-delà" in lu.motif


def test_le_texte_d_un_controle_de_contenu_est_lu() -> None:
    corps = (
        "<w:sdt><w:sdtPr/><w:sdtContent><w:p><w:r><w:t>Marché régional : "
        "14,2 M€ (CCI 2025)</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    )
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        z.writestr(
            "word/document.xml", f"<w:document {_W}><w:body>{corps}</w:body></w:document>"
        )
    assert "14,2 M€ (CCI 2025)" in dc.extraire("modele.docx", tampon.getvalue()).texte


def test_les_pages_non_lues_font_un_pdf_lu_en_partie(monkeypatch: pytest.MonkeyPatch) -> None:
    from reportlab.pdfgen import canvas

    tampon = BytesIO()
    toile = canvas.Canvas(tampon)
    for page in range(2):
        for rang in range(4):
            toile.drawString(72, 760 - 14 * rang, f"Page {page} : texte suffisant pour etre lu.")
        toile.showPage()
    toile.save()
    monkeypatch.setattr(dc, "MAX_PAGES_PDF", 1)

    lu = dc.extraire("rapport.pdf", tampon.getvalue())

    assert lu.statut == StatutLecture.TRONQUE
    assert "pages 2 à 2 non lues" in lu.motif
    assert "[page" not in lu.texte, "aucun repère de lecture dans le texte transmis"


def test_une_relance_apres_un_chapitre_ecrit_ne_lit_pas_la_bibliotheque(
    job_avec_organisation: Any,
) -> None:
    """S4 : aucun document au premier passage, un dépôt, une relance — rien de lu.

    Le socle, déjà verrouillé, n'aurait jamais vu ce document ; les chapitres
    restants l'auraient reçu.
    """
    job = job_avec_organisation
    assert dc.lire_les_documents(job) == []
    chapitre = job.chapters.first()
    chapitre.status = "done"
    chapitre.save(update_fields=["status"])
    _deposer(job, "tardif.docx", _docx(["Déposé après le lancement."]))

    assert dc.lire_les_documents(job) == []
    assert dc.bloc_documents(job) == ""


def test_retirer_un_document_pendant_la_generation_ne_vide_pas_ses_chapitres(
    job_avec_organisation: Any,
) -> None:
    """S8 : le texte attend la fin du dossier pour partir."""
    from generation.models import JobStatus

    job = job_avec_organisation
    piece = _deposer(job, "bilan.docx", _docx(["Résultat net : 12 400 €."]))
    dc.lire_les_documents(job)
    job.status = JobStatus.RUNNING
    job.save(update_fields=["status"])

    piece.delete()
    assert "12 400" in dc.bloc_documents(job), "les chapitres restants le lisent encore"

    job.status = JobStatus.DONE
    job.save(update_fields=["status"])
    dc.effacer_les_textes_orphelins()
    assert dc.bloc_documents(job) == ""
    assert DocumentClientLu.objects.get(job=job).texte_efface_le is not None


def test_supprimer_l_organisation_efface_aussi_le_texte(db: Any) -> None:
    """Chemin CASCADE : l'organisation part, ses pièces et leur texte avec elle."""
    from catalog.models import DeliverableType, Offer
    from customers.models import Customer
    from generation.services import bootstrap_generation_job
    from intake.models import IntakeStatus, IntakeSubmission
    from orders.models import Order
    from organisations import services
    from organisations.models import PieceJointe

    contact = Customer.objects.create(email="cascade@exemple.fr")
    organisation = services.creer_organisation(raison_sociale="Cascade", contact=contact)
    offre = Offer.objects.create(
        name="S", slug="s-cascade", deliverable_type=DeliverableType.BUSINESS_STRATEGY
    )
    commande = Order.objects.create(
        systeme_order_id="c-cascade", customer=contact, offer=offre
    )
    soumission = IntakeSubmission.objects.create(
        order=commande, status=IntakeStatus.NORMALIZED, normalized_variables={"SECTEUR": "x"},
    )
    job = bootstrap_generation_job(soumission)
    piece = PieceJointe(organisation=organisation, categorie="document", nom_original="b.docx")
    piece.fichier.save("b.docx", ContentFile(_docx(["Contenu du bilan."])), save=False)
    piece.save()
    job.status = "done"
    job.save(update_fields=["status"])
    DocumentClientLu.objects.create(
        job=job, piece=piece, nom="b.docx", statut=StatutLecture.LU, texte="Contenu du bilan."
    )

    organisation.delete()

    assert DocumentClientLu.objects.get(job=job).texte == ""


def test_seuls_les_montants_avec_unite_des_documents_sont_admis(
    job_avec_organisation: Any,
) -> None:
    """S9 : les nombres nus d'un prévisionnel ne justifient pas un chiffre inventé."""
    from generation.verification.services import chiffres_du_brief

    job = job_avec_organisation
    _deposer(job, "p.docx", _docx(["Chiffre d'affaires : 45 000 €. Foyers : 38 000."]))
    dc.lire_les_documents(job)

    admis = chiffres_du_brief(job)
    assert 45000.0 in admis
    assert 38000.0 not in admis


def test_une_reprise_lit_la_bibliotheque_du_dossier_d_origine(
    client_admin: Any, job_avec_organisation: Any, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """S10 : pas celle d'une autre organisation dont le client serait contact."""
    import json

    from generation import tasks
    from generation.models import GenerationJob
    from organisations import services

    origine = job_avec_organisation
    autre = services.creer_organisation(
        raison_sociale="Autre maison", contact=origine.order.customer
    )
    monkeypatch.setattr(tasks.run_generation_job_task, "delay", lambda *a, **k: None)
    reponse = client_admin.post(f"/api/dashboard/jobs/{origine.id}/regenerer/")
    reprise = GenerationJob.objects.get(id=json.loads(reponse.content)["job_id"])

    assert dc.organisation_des_documents(reprise) == origine.order.organisation
    assert dc.organisation_des_documents(reprise) != autre


def test_sous_l_ancien_moteur_rien_n_est_marque_lu(
    job_avec_organisation: Any, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """S11 : aucun prompt de l'ancien moteur ne reçoit ce texte."""
    from generation import runner as moteur

    class _Arret(Exception):  # noqa: N818
        pass

    def _arreter(*args: Any, **kwargs: Any) -> Any:
        raise _Arret

    job = job_avec_organisation
    _deposer(job, "etude.docx", _docx(["Marché régional : 14,2 M€."]))
    monkeypatch.setattr(moteur, "debiter_pour_job", lambda job: (True, ""))
    monkeypatch.setattr(moteur, "_moteur_structure", lambda job: False)
    monkeypatch.setattr(moteur, "_build_phase0_plan", _arreter)

    with pytest.raises(_Arret):
        moteur.run_generation_job(job)

    assert job.documents_client.count() == 0


def test_les_intitules_du_prompt_sont_des_labels_internes() -> None:
    from generation.internal_labels import INTERNAL_LABEL_NAMES

    assert "DOCUMENTS_DU_CLIENT" in INTERNAL_LABEL_NAMES
    assert "BRIEF_CLIENT" in INTERNAL_LABEL_NAMES


def test_une_reprise_sans_envoi_passe_le_verrou_du_document(
    job_avec_organisation: Any, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """S7 : un document amputé ne devient ni téléchargeable ni envoyable."""
    from delivery import services as livraison
    from generation.models import GenerationJob, JobStatus
    from monitoring.models import OperationalIncident

    job = job_avec_organisation
    job.status = JobStatus.DONE
    job.save(update_fields=["status"])
    monkeypatch.setattr(livraison, "get_pdf_client", lambda: None)
    monkeypatch.setattr(
        livraison, "_assembler_livrable",
        lambda j, pdf_client: livraison.Assemblage(
            artefacts=(), url_principale="", retenu="tableau du compte de résultat vide"
        ),
    )

    livraison.assembler_sans_envoyer(job)

    assert GenerationJob.objects.get(pk=job.pk).status == JobStatus.INTERVENTION_REQUISE
    assert OperationalIncident.objects.filter(job=job, title__contains="retenu").exists()


# ── Contre-relecture du 11/09/2026 ───────────────────────────────────────────


def test_deux_paragraphes_d_une_cellule_ne_se_collent_pas() -> None:
    """« CA 2025 » / « 45 000 € » dans une cellule se lisait « CA 202545 000 € »."""
    import docx

    from core.numbers import amounts_in

    document = docx.Document()
    cellule = document.add_table(rows=1, cols=2).rows[0].cells[1]
    cellule.text = "CA 2025"
    cellule.add_paragraph("45 000 €")
    tampon = BytesIO()
    document.save(tampon)

    lu = dc.extraire("prev.docx", tampon.getvalue())

    assert "CA 2025 / 45 000 €" in lu.texte, lu.texte
    assert 202545000.0 not in amounts_in(lu.texte)


def test_deux_paragraphes_d_une_cellule_de_diapositive_ne_se_collent_pas() -> None:
    a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    cellule = (
        "<a:tc><a:txBody><a:p><a:r><a:t>2025</a:t></a:r></a:p>"
        "<a:p><a:r><a:t>45 000</a:t></a:r></a:p></a:txBody></a:tc>"
    )
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        z.writestr(
            "ppt/slides/slide1.xml",
            f'<p:sld {a} xmlns:p="p"><a:tbl><a:tr>{cellule}</a:tr></a:tbl></p:sld>',
        )
    assert "2025 / 45 000" in dc.extraire("d.pptx", tampon.getvalue()).texte


def _classeur(lignes: list[str], formats: dict[int, str], styles: list[int]) -> bytes:
    """Un classeur d'une feuille « Prévisionnel ». `formats` : id → code."""
    numfmts = "".join(
        f'<numFmt numFmtId="{i}" formatCode="{code}"/>' for i, code in formats.items()
    )
    xfs = "".join(f'<xf numFmtId="{i}"/>' for i in styles)
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        z.writestr("xl/workbook.xml", (
            f"<workbook {_S} {_R}><sheets>"
            '<sheet name="Prévisionnel" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ))
        z.writestr("xl/_rels/workbook.xml.rels", (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="x" Target="worksheets/sheet1.xml"/></Relationships>'
        ))
        z.writestr("xl/styles.xml", (
            f"<styleSheet {_S}><numFmts>{numfmts}</numFmts>"
            f"<cellXfs>{xfs}</cellXfs></styleSheet>"
        ))
        z.writestr(
            "xl/worksheets/sheet1.xml",
            f"<worksheet {_S}><sheetData>{''.join(lignes)}</sheetData></worksheet>",
        )
    return tampon.getvalue()


def _rangee(numero: int, *cellules: tuple[str, str | None]) -> str:
    """`(valeur, style)` ; un style None avec du texte fait une chaîne en ligne."""
    xml = ""
    for rang, (valeur, style) in enumerate(cellules):
        ref = f"{'ABCD'[rang]}{numero}"
        if style is None:
            xml += f'<c r="{ref}" t="inlineStr"><is><t>{valeur}</t></is></c>'
        else:
            xml += f'<c r="{ref}" s="{style}"><v>{valeur}</v></c>'
    return f'<row r="{numero}">{xml}</row>'


def test_un_montant_au_format_monetaire_garde_son_symbole() -> None:
    """N1 : `#,##0 "€"` rendait `45000` — que le contrôle ne reconnaissait plus."""
    classeur = _classeur(
        [_rangee(1, ("CA", None), ("45000", "1"), ("1234.5", "2"))],
        {164: '#,##0 &quot;€&quot;', 165: '[$€-40C] #,##0.00'},
        [0, 164, 165],
    )
    texte = dc.extraire("p.xlsx", classeur).texte
    assert "CA | 45 000 € | 1 234,5 €" in texte, texte


def test_la_repetition_joue_le_texte_que_l_extraction_produit_vraiment() -> None:
    """N1 : la répétition prétendait jouer « le texte tel que l'extraction le
    rend », mais son texte portait « 45 000 € » quand l'extraction rendait
    « 45000 ». Le voici produit par l'extraction elle-même, à partir d'un
    classeur réel : s'ils divergent, la répétition joue une entrée idéale.
    """
    from generation.repetition import DOCUMENT_DE_REPETITION

    premier_septembre = (date(2026, 9, 1) - date(1899, 12, 30)).days
    classeur = _classeur(
        [
            _rangee(1, ("Poste", None), ("2026", "0"), ("2027", "0"), ("2028", "0")),
            _rangee(2, ("Chiffre d'affaires", None), ("45000", "1"), ("62000", "1"),
                    ("80000", "1")),
            _rangee(3, ("Marge brute", None), ("0.38", "2"), ("0.4", "2"), ("0.41", "2")),
            _rangee(4, ("Date de mise à jour", None), (str(premier_septembre), "3")),
        ],
        {164: '#,##0 &quot;€&quot;'},
        [0, 164, 9, 14],
    )
    assert dc.extraire("previsionnel.xlsx", classeur).texte == DOCUMENT_DE_REPETITION


def test_une_reprise_de_reprise_remonte_au_dossier_d_origine(
    client_admin: Any, job_avec_organisation: Any, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import json

    from generation import tasks
    from generation.models import GenerationJob
    from organisations import services

    origine = job_avec_organisation
    services.creer_organisation(raison_sociale="Autre maison", contact=origine.order.customer)
    monkeypatch.setattr(tasks.run_generation_job_task, "delay", lambda *a, **k: None)
    premiere = GenerationJob.objects.get(id=json.loads(
        client_admin.post(f"/api/dashboard/jobs/{origine.id}/regenerer/").content
    )["job_id"])
    seconde = GenerationJob.objects.get(id=json.loads(
        client_admin.post(f"/api/dashboard/jobs/{premiere.id}/regenerer/").content
    )["job_id"])

    assert dc.organisation_des_documents(seconde) == origine.order.organisation


def test_un_socle_en_echec_n_empeche_pas_la_lecture(job_avec_organisation: Any) -> None:
    """Un socle refusé est reconstruit de zéro à la relance : il doit voir les documents."""
    from generation.models import SocleDonnees, SocleStatut

    job = job_avec_organisation
    SocleDonnees.objects.create(job=job, statut=SocleStatut.INVALIDE, contenu={})
    _deposer(job, "etude.docx", _docx(["Marché régional : 14,2 M€."]))

    assert len(dc.lire_les_documents(job)) == 1


def test_une_correction_en_cours_garde_la_matiere_de_ses_chapitres(
    job_avec_organisation: Any,
) -> None:
    """S8 : la boucle de correction réécrit des chapitres sur un dossier TERMINÉ."""
    from generation.models import JobStatus

    job = job_avec_organisation
    piece = _deposer(job, "bilan.docx", _docx(["Résultat net : 12 400 €."]))
    dc.lire_les_documents(job)
    job.status = JobStatus.DONE
    job.save(update_fields=["status"])
    chapitre = job.chapters.first()
    chapitre.status = "running"
    chapitre.save(update_fields=["status"])

    piece.delete()
    dc.effacer_les_textes_orphelins()
    assert "12 400" in dc.bloc_documents(job), "le chapitre en réécriture le lit encore"

    chapitre.status = "done"
    chapitre.save(update_fields=["status"])
    dc.effacer_les_textes_orphelins()
    assert dc.bloc_documents(job) == ""


def test_les_cellules_sans_valeur_se_comptent_pour_tout_le_classeur() -> None:
    """Prévisionnel réel du 11/09/2026 : « 3 cellules… ; 51 cellules… », une par feuille."""
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        z.writestr("xl/workbook.xml", (
            f"<workbook {_S} {_R}><sheets>"
            '<sheet name="A" sheetId="1" r:id="rId1"/><sheet name="B" sheetId="2" r:id="rId2"/>'
            "</sheets></workbook>"
        ))
        z.writestr("xl/_rels/workbook.xml.rels", (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="x" Target="worksheets/sheet1.xml"/>'
            '<Relationship Id="rId2" Type="x" Target="worksheets/sheet2.xml"/>'
            "</Relationships>"
        ))
        for numero in (1, 2):
            z.writestr(f"xl/worksheets/sheet{numero}.xml", (
                f"<worksheet {_S}><sheetData><row r=\"1\">"
                '<c r="A1" t="inlineStr"><is><t>Total</t></is></c>'
                '<c r="B1"><f>SUM(A1)</f></c></row></sheetData></worksheet>'
            ))
    lu = dc.extraire("p.xlsx", tampon.getvalue())
    assert lu.motif.count("cellule(s) calculée(s)") == 1, lu.motif
    assert "2 cellule(s) calculée(s)" in lu.motif
