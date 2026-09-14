"""Trois chiffres que le lecteur voit justifiés, et que le contrôle accusait.

## Relevés dans les 233 « chiffres hors socle » des business plans et stratégies
## (corpus du 14/09/2026)

- Les tableaux GO / AJUSTER / STOP que le manuel impose : « Seuil STOP :
  … | Toujours sous -20 % », « Seuil ADJUST : … | Entre 100 000 € et
  175 000 € engagés ». Sous un en-tête de seuil, la valeur est une DÉCISION du
  projet, pas un fait de marché.
- Le choc d'un scénario de sensibilité : « Scénario dégradé (-30 %) »,
  « absorbe un choc de chiffre d'affaires jusqu'à -66 % ». C'est le paramètre
  qu'on applique.
- « Scénario dégradé : CA année 1 | 54 276 € | 37 993 € | 54 276 € moins
  30 %, soit 54 276 € × 0,70 = 37 993 € » : la valeur est posée par un calcul
  dans la colonne SUIVANTE, et n'était jugée qu'à sa première occurrence.

Contre-épreuves : un taux sous un en-tête qui ne décide rien, une croissance
hors scénario, un calcul faux posé plus loin restent signalés.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from generation.socle.referentiel import Fiabilite, Perimetre
from generation.socle.schema import DonneeSocle, Socle, Zone
from generation.verification.controles import controler_chiffres_hors_socle
from generation.verification.lecture import lire_livrable


def _socle() -> Socle:
    return Socle(
        secteur="services", zone=Zone(pays="France"), date_socle=date(2026, 9, 14),
        donnees=[DonneeSocle(
            id="ca_previsionnel_an1", libelle="Chiffre d'affaires prévisionnel — exercice 1",
            valeur=54276, unite="EUR", annee=2026, perimetre=Perimetre.ENTREPRISE,
            fiabilite=Fiabilite.DECLAREE, source="données du projet",
        )],
    )


def _signales(
    tmp_path: Path, paragraphes: list[str], tableau: list[list[str]] | None = None,
) -> set[str]:
    from docx import Document

    document = Document()
    for texte in paragraphes:
        document.add_paragraph(texte)
    if tableau:
        table = document.add_table(rows=len(tableau), cols=len(tableau[0]))
        for rang, ligne in enumerate(tableau):
            for colonne, valeur in enumerate(ligne):
                table.cell(rang, colonne).text = valeur
    chemin = tmp_path / "bp.docx"
    document.save(str(chemin))
    anomalies = controler_chiffres_hors_socle(lire_livrable(chemin), _socle())
    return {a.detail.split(" »")[0].lstrip("« ") for a in anomalies}


def test_un_seuil_de_decision_est_une_decision(tmp_path: Path) -> None:
    signales = _signales(tmp_path, ["Pilotage."], [
        ["Décision", "Seuil GO", "Seuil ADJUST", "Seuil STOP"],
        ["Levée de fonds", "175 000 € contractualisés", "Entre 100 000 € et 175 000 €",
         "Toujours sous -20 %"],
    ])
    assert not signales & {"175 000 €", "100 000 €", "-20 %"}, signales


def test_un_en_tete_qui_ne_decide_rien_ne_blanchit_rien(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : une part, et un seuil de RENTABILITÉ — qui est un calcul."""
    signales = _signales(tmp_path, ["Pilotage."], [
        ["Poste", "Part du chiffre d'affaires", "Seuil de rentabilité"],
        ["Abonnements", "47,1 %", "19 500 €"],
    ])
    assert {"47,1 %", "19 500 €"} <= signales


def test_le_choc_d_un_scenario_est_un_parametre(tmp_path: Path) -> None:
    signales = _signales(tmp_path, [
        "Dans le scénario dégradé, le chiffre d'affaires recule de 30 %.",
        "Le modèle absorbe un choc de chiffre d'affaires jusqu'à -66 % environ.",
    ])
    assert not signales & {"30 %", "-66 %"}, signales


def test_une_croissance_hors_scenario_reste_un_fait(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : « scénario » seul ne blanchit pas une croissance de marché."""
    signales = _signales(tmp_path, ["Le marché croît de 6,2 % par an."])
    assert "6,2 %" in signales


def test_une_valeur_posee_par_un_calcul_plus_loin_dans_sa_ligne(tmp_path: Path) -> None:
    signales = _signales(tmp_path, ["Scénarios."], [
        ["Scénario dégradé", "Central", "Dégradé", "Calcul"],
        ["Chiffre d'affaires année 1", "54 276 €", "37 993 €",
         "54 276 € moins 30 %, soit 54 276 € × 0,70 = 37 993 €"],
    ])
    assert "37 993 €" not in signales


def test_un_calcul_faux_pose_plus_loin_reste_signale(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : 54 276 € × 0,70 ne fait pas 39 993 €."""
    signales = _signales(tmp_path, ["Scénarios."], [
        ["Scénario dégradé", "Central", "Dégradé", "Calcul"],
        ["Chiffre d'affaires année 1", "54 276 €", "39 993 €",
         "54 276 € moins 30 %, soit 54 276 € × 0,70 = 39 993 €"],
    ])
    assert "39 993 €" in signales
