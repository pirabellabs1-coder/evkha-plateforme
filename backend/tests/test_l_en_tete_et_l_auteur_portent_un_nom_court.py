"""L'en-tête tient sur une ligne, et l'auteur du PDF est le porteur du projet.

29/09/2026, business plan ÉCLORE (« Et deux défauts visibles qui n'étaient pas
dans la liste », diagnostic) : l'en-tête courant ET l'auteur du PDF prenaient la
raison sociale saisie dans « Ma marque » — ici la phrase entière « ÉCLORE (nom
de projet provisoire), avec pour signature « Expériences bien-être… » » — sur
105 pages. Le nom du porteur était collecté (`PORTEUR_PROJET`) et jamais lu par
le rendu.

Tenu ici, sur le `.docx` écrit (règle 7) : l'en-tête porte un nom COURT —
celui du projet s'il en est un, sinon la tête de la raison sociale, coupée au
mot —, et l'auteur est le porteur, à défaut le nom court de la marque. Le
dernier test traverse la base : sans lui, un rendu juste resterait sourd aux
variables du dossier.
"""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from generation.rendu_word.depuis_json import rendre_etude
from generation.rendu_word.texte import NOM_COURT_MAX, nom_court

RAISON_SOCIALE_ECLORE = (
    "ÉCLORE (nom de projet provisoire), avec pour signature « Expériences "
    "bien-être en pleine nature, pour se retrouver »"
)


def _rendre(tmp_path: Path, **etude: Any) -> Path:
    base: dict[str, Any] = {
        "titre": "Business plan",
        "marque": {"nom": RAISON_SOCIALE_ECLORE},
        "chapitres": [],
    }
    base.update(etude)
    return rendre_etude(base, tmp_path / "document.docx")


def _entete(chemin: Path) -> str:
    return " | ".join(
        p.text for section in Document(str(chemin)).sections
        for p in section.header.paragraphs if p.text
    )


def _auteur(chemin: Path) -> str:
    return str(Document(str(chemin)).core_properties.author)


# ── L'en-tête ────────────────────────────────────────────────────────────────


def test_l_en_tete_ne_recopie_pas_la_raison_sociale_entiere(tmp_path: Path) -> None:
    """Le défaut exact d'ÉCLORE : la phrase de « Ma marque » sur 105 pages."""
    assert _entete(_rendre(tmp_path)) == "ÉCLORE  /  Business plan"


def test_l_en_tete_prefere_le_nom_du_projet(tmp_path: Path) -> None:
    chemin = _rendre(tmp_path, projet="Éclore Nature")
    assert _entete(chemin) == "Éclore Nature  /  Business plan"


def test_une_description_de_projet_n_est_pas_un_nom(tmp_path: Path) -> None:
    """`PROJET` porte parfois une description : on ne l'imprime pas coupée."""
    chemin = _rendre(
        tmp_path,
        projet="Atelier de torréfaction avec vente directe, abonnements et formations",
        marque={"nom": "Maison Lorel"},
    )
    assert _entete(chemin) == "Maison Lorel  /  Business plan"


def test_un_nom_deja_court_traverse_intact(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : « Saint-Étienne » et « L'Atelier » ne se coupent pas."""
    chemin = _rendre(tmp_path, marque={"nom": "L'Atelier de Saint-Étienne"})
    assert _entete(chemin) == "L'Atelier de Saint-Étienne  /  Business plan"


@pytest.mark.parametrize(
    ("brut", "attendu"),
    [
        (RAISON_SOCIALE_ECLORE, "ÉCLORE"),
        ("ÉCLORE, expériences bien-être", "ÉCLORE"),
        ("ÉCLORE — expériences bien-être", "ÉCLORE"),
        ("ÉCLORE avec pour signature « Expériences bien-être »", "ÉCLORE"),
        ("« ÉCLORE »", "ÉCLORE"),
        ("ÉCLORE - expériences", "ÉCLORE"),
        ("Clémence Martin, fondatrice", "Clémence Martin"),
        ("Maison Lorel", "Maison Lorel"),
        ("", ""),
    ],
)
def test_le_nom_court_s_arrete_a_la_premiere_apposition(brut: str, attendu: str) -> None:
    """La classe des appositions, pas le seul cas vu (règle 4)."""
    assert nom_court(brut) == attendu


@pytest.mark.parametrize(
    "raison_sociale",
    [
        "Martin, Durand & Associés",
        "SAS « Les Délices »",
        "Dupont & Fils S.A.",
        "Martin, Clémence",
        "ÉCLORE « Expériences bien-être »",
    ],
)
def test_une_raison_sociale_legitime_n_est_pas_mutilee(raison_sociale: str) -> None:
    """Revue du 29/09/2026 : « Martin », « SAS », « Dupont & Fils S.A ».

    Une virgule suivie d'une CAPITALE joint des noms ; un guillemet et un point
    appartiennent au nom. Seule une apposition OUVERTE — parenthèse, tiret long,
    « avec pour », virgule suivie d'une minuscule — le termine.
    """
    assert nom_court(raison_sociale) == raison_sociale


def test_une_description_courte_ne_passe_pas_pour_un_nom(tmp_path: Path) -> None:
    """Revue du 29/09/2026 : le formulaire Tally range la DESCRIPTION dans `PROJET`.

    « Salon de coiffure mixte, situé à Lyon… » donnait l'en-tête « Salon de
    coiffure mixte », imprimé sur chaque page, au lieu de la marque.
    """
    chemin = _rendre(
        tmp_path,
        projet="Salon de coiffure mixte, situé à Lyon, proposant coupes et colorations",
        marque={"nom": "Maison Lorel"},
    )
    assert _entete(chemin) == "Maison Lorel  /  Business plan"


@pytest.mark.parametrize(
    "projet", ["Boulangerie du Parc", "L'Atelier de Saint-Étienne", "ÉCLORE", "Éclore Nature"],
)
def test_un_nom_de_projet_reste_prefere_a_la_marque(tmp_path: Path, projet: str) -> None:
    """CONTRE-ÉPREUVE : un vrai nom de projet l'emporte toujours sur la marque."""
    chemin = _rendre(tmp_path, projet=projet, marque={"nom": "Maison Lorel"})
    assert _entete(chemin) == f"{projet}  /  Business plan"


def test_un_nom_sans_apposition_trop_long_se_coupe_au_mot() -> None:
    long = "Compagnie générale des expériences de bien-être en pleine nature du Vercors"
    court = nom_court(long)
    assert len(court) <= NOM_COURT_MAX
    assert court.endswith("…")
    assert long.startswith(court[:-1]) and long[len(court) - 1] == " ", court


# ── L'auteur du fichier ──────────────────────────────────────────────────────


def test_l_auteur_est_le_porteur_du_projet(tmp_path: Path) -> None:
    """Un business plan est présenté par son porteur."""
    assert _auteur(_rendre(tmp_path, porteur="Clémence Martin")) == "Clémence Martin"


def test_sans_porteur_l_auteur_est_le_nom_court_de_la_marque(tmp_path: Path) -> None:
    """Le défaut exact d'ÉCLORE : la phrase entière en auteur du PDF."""
    assert _auteur(_rendre(tmp_path)) == "ÉCLORE"


def test_le_titre_du_fichier_reste_celui_du_document(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : seul l'auteur change."""
    chemin = _rendre(tmp_path, porteur="Clémence Martin")
    assert Document(str(chemin)).core_properties.title == "Business plan"


def test_rien_n_est_invente_sans_nom(tmp_path: Path) -> None:
    """CONTRE-ÉPREUVE : sans porteur ni marque, l'auteur reste vide."""
    assert _auteur(_rendre(tmp_path, marque={})) == ""


# ── Depuis la base : les variables du dossier atteignent le rendu ────────────


@pytest.mark.django_db
def test_le_dossier_reel_porte_son_projet_et_son_porteur(tmp_path: Path) -> None:
    """`PROJET` et `PORTEUR_PROJET` étaient collectés et jamais lus."""
    from catalog.models import DeliverableType, Offer
    from customers.models import Customer
    from generation.models import (
        ChapterGeneration,
        ChapterStatus,
        SocleDonnees,
        SocleStatut,
    )
    from generation.rendu_word.services import produire_docx
    from generation.services import bootstrap_generation_job
    from intake.models import IntakeStatus, IntakeSubmission
    from orders.models import Order
    from tests.test_lot3_livrable import _chapitre, _socle

    offre = Offer.objects.create(
        name="Business plan", slug="bp-nom-court",
        deliverable_type=DeliverableType.BUSINESS_PLAN,
    )
    client = Customer.objects.create(email="nom-court@example.com")
    commande = Order.objects.create(
        systeme_order_id="order-nom-court", customer=client, offer=offre,
    )
    soumission = IntakeSubmission.objects.create(
        order=commande, status=IntakeStatus.NORMALIZED,
        normalized_variables={
            "SECTEUR": "bien-être", "PAYS": "France", "ZONE": "Drôme",
            "PROJET": "ÉCLORE", "PORTEUR_PROJET": "Clémence Martin",
            "NOM_ENTREPRISE": RAISON_SOCIALE_ECLORE,
        },
    )
    job = bootstrap_generation_job(soumission)
    SocleDonnees.objects.create(
        job=job, statut=SocleStatut.VALIDE, contenu=_socle().model_dump(mode="json"),
    )
    ChapterGeneration.objects.filter(job=job, chapter_number=1).update(
        status=ChapterStatus.DONE, payload=_chapitre(1).model_dump(mode="json"),
    )

    chemin = produire_docx(job, destination=tmp_path / "reel.docx").chemin

    assert _entete(chemin) == "ÉCLORE  /  Business plan"
    with zipfile.ZipFile(chemin) as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
    assert "<dc:creator>Clémence Martin</dc:creator>" in core, core
    assert "nom de projet provisoire" not in core
