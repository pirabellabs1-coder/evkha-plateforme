"""Aucun mot coupé en deux dans un tableau du document livré.

Relevé par la cliente le 29/10/2026, sur les deux études de marché de la
boutique : « des troncatures au niveau de certains tableaux ». Le PDF de
`99a142b0` (rendu LibreOffice, production) montrait « recommandati-ons »,
« automatisatio-n », « documentatio-n » : des mots coupés lettre à lettre dans
des colonnes plus étroites qu'eux.

Deux causes, mesurées sur le corpus de production (20 documents, 1 226
tableaux de données) :

1. la GRILLE du tableau (`w:gridCol`) restait à parts égales alors que les
   cellules avaient des largeurs proportionnelles. LibreOffice, qui fabrique le
   PDF livré, lit la grille : le correctif « tableaux trop étroits » du 09/08
   n'avait jamais atteint le PDF ;
2. même proportionnelle, une colonne pouvait être plus étroite que son mot le
   plus long — 96 tableaux du corpus sur 1 226 ; 0 après correctif, dont 2
   coupés en deux tableaux faute de place à 7 pt.
"""
from __future__ import annotations

import itertools
import random

import pytest
from docx.oxml.ns import qn
from docx.table import Table

from generation.rendu_word import composants
from generation.rendu_word.composants import (
    FORME_TABLEAU_CONFORTABLE,
    LARGEUR_UTILE_DXA,
    _largeurs,
    _planchers,
    _relever_aux_planchers,
    forme_du_tableau,
)
from generation.rendu_word.gabarit import charger_gabarit
from generation.rendu_word.palette import REF_CREME, REF_PRUNE, construire_palette

#: Le tableau à huit colonnes de l'étude IA `99a142b0`, chapitre 19, tel que
#: livré le 28/10/2026.
ENTETES_99A = [
    "Cible", "Zone", "Offre EVKHA", "Canal", "Partenaire", "Priorité", "Délai",
    "Indicateur de réussite",
]
LIGNES_99A = [
    ["Artisans du BTP débordés", "Ville moyenne et périphérie",
     "Audit express devis-facture", "Rendez-vous en présentiel",
     "CMA, réseau d'artisans", "1", "0 à 90 jours",
     "2 missions signées sur 6 audits"],
    ["Artisans clients d'un cabinet comptable", "Ville moyenne",
     "Mise en place facture et relances", "Recommandation",
     "Expert-comptable", "1", "0 à 6 mois",
     "1 mission par mois issue du cabinet"],
]


def _rendre(entetes: list[str], lignes: list[list[str]], **options: str) -> list[Table]:
    document = charger_gabarit()
    palette = construire_palette(primaire=REF_PRUNE, fond_clair=REF_CREME)
    avant = len(document.tables)
    composants.tableau(document, palette, entetes, lignes, **options)
    return list(document.tables[avant:])


def _grille(table: Table) -> list[int]:
    return [
        int(colonne.get(qn("w:w")))
        for colonne in table._tbl.tblGrid.findall(qn("w:gridCol"))
    ]


def _cellules(table: Table) -> list[int]:
    return [
        int(cellule._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w")))
        for cellule in table.rows[0].cells
    ]


# ── La grille ────────────────────────────────────────────────────────────────


def test_la_grille_porte_les_memes_largeurs_que_les_cellules() -> None:
    """Le défaut du PDF : grille égale, cellules proportionnelles.

    Sur l'ancien code, la grille valait 8 × 1 204 DXA quel que soit le contenu.
    """
    (table,) = _rendre(ENTETES_99A, LIGNES_99A)

    assert _grille(table) == _cellules(table)
    assert len(set(_grille(table))) > 1


def test_l_encadre_garde_son_rendu() -> None:
    """CONTRE-ÉPREUVE (relecture du 30/10) : aligner la grille de l'encadré sur
    ses largeurs fixes réduisait l'étiquette à 1 400 DXA de texte, sous
    « RECOMMANDATION » (2 396) — le défaut corrigé, déplacé. Hors tableau de
    données, la grille et les marges restent celles d'avant."""
    document = charger_gabarit()
    palette = construire_palette(primaire=REF_PRUNE, fond_clair=REF_CREME)
    composants.encadre(document, palette, "Recommandation", ["Ligne"], verdict=False)
    table = document.tables[-1]

    assert len(set(_grille(table))) == 1
    assert len(table._tbl.tblPr.findall(qn("w:tblCellMar"))) == 2


@pytest.mark.parametrize(
    "cellule",
    ["1 250 000 €", "2026-2030", "2025/12"],
)
def test_un_bloc_insecable_se_mesure_en_entier(cellule: str) -> None:
    """Le moteur ne coupe ni à l'insécable d'un montant, ni entre un trait
    d'union et un chiffre : le plancher compte le bloc entier."""
    forme = FORME_TABLEAU_CONFORTABLE
    (plancher,) = _planchers([""], [[cellule]], 1, forme)
    marges = 2 * forme.marge

    assert plancher - marges > _planchers([""], [["000"]], 1, forme)[0] - marges + 300


def test_l_arrondi_ne_fait_passer_aucune_colonne_sous_son_plancher() -> None:
    """Le cas trouvé par la relecture : 1 556 pour un plancher de 1 557."""
    largeurs = [1571, 904, 954, 1191, 773, 1219, 1608, 1415]
    planchers = [1557, 691, 365, 1012, 548, 2500, 2347, 614]

    resultat = _relever_aux_planchers(largeurs, planchers)

    assert sum(resultat) == LARGEUR_UTILE_DXA
    assert all(r >= p for r, p in zip(resultat, planchers, strict=True)), resultat


def test_les_marges_interieures_ne_sont_posees_qu_une_fois() -> None:
    """`tableau` redéfinissait les marges après `_table` : deux `tblCellMar`."""
    (table,) = _rendre(ENTETES_99A, LIGNES_99A)

    assert len(table._tbl.tblPr.findall(qn("w:tblCellMar"))) == 1


# ── Aucun mot coupé ──────────────────────────────────────────────────────────


def test_le_tableau_livre_de_99a142b0_ne_coupe_plus_aucun_mot() -> None:
    """Le cas exact : « Recommandation », « Indicateur », « Partenaire »."""
    forme = forme_du_tableau(ENTETES_99A, LIGNES_99A)
    largeurs = _largeurs(ENTETES_99A, LIGNES_99A, 8, forme)
    planchers = _planchers(ENTETES_99A, LIGNES_99A, 8, forme)

    assert all(
        largeur >= plancher
        for largeur, plancher in zip(largeurs, planchers, strict=True)
    ), list(zip(ENTETES_99A, largeurs, planchers, strict=True))


def test_a_parts_egales_ce_tableau_coupait_des_mots() -> None:
    """La grille livrée le 28/10 : huit parts égales, plus étroites que
    « Recommandation » ou « Indicateur » à 9 pt."""
    planchers = _planchers(ENTETES_99A, LIGNES_99A, 8, FORME_TABLEAU_CONFORTABLE)

    assert any(plancher > LARGEUR_UTILE_DXA // 8 for plancher in planchers)


def test_un_mot_sans_fin_ne_vide_pas_la_colonne_voisine() -> None:
    """Un mot plus large que le plafond sera coupé de toute façon : il ne
    réclame pas plus que ce plafond."""
    entetes = ["Oui", "Référence"]
    lignes = [["Oui", "x" * 400]]
    forme = forme_du_tableau(entetes, lignes)

    largeurs = _largeurs(entetes, lignes, 2, forme)

    assert largeurs[0] >= _planchers(entetes, lignes, 2, forme)[0]


def test_les_largeurs_font_toujours_exactement_la_page_et_respectent_les_planchers() -> None:
    """Propriété, sur des tableaux tirés au hasard — dont l'arrondi de 1 DXA
    qui faisait passer une colonne sous son plancher (`d4d73208`)."""
    tirage = random.Random(30_10_2026)
    mots = [
        "a", "Oui", "Haute", "Priorité", "Partenaire", "Recommandation",
        "accompagnement", "2 000 €", "expert-comptable", "https://numeum.fr/a/b",
    ]
    for _ in range(400):
        colonnes = tirage.randint(2, 7)
        entetes = [tirage.choice(mots) for _ in range(colonnes)]
        lignes = [
            [" ".join(tirage.choices(mots, k=tirage.randint(1, 6))) for _ in range(colonnes)]
            for _ in range(tirage.randint(1, 5))
        ]
        forme = forme_du_tableau(entetes, lignes)
        largeurs = _largeurs(entetes, lignes, colonnes, forme)
        planchers = _planchers(entetes, lignes, colonnes, forme)

        assert sum(largeurs) == LARGEUR_UTILE_DXA, (entetes, largeurs)
        if sum(planchers) <= LARGEUR_UTILE_DXA:
            assert all(
                largeur >= plancher
                for largeur, plancher in zip(largeurs, planchers, strict=True)
            ), (entetes, lignes, largeurs, planchers)


def test_un_tableau_ordinaire_garde_son_corps_et_ses_marges() -> None:
    """CONTRE-ÉPREUVE : 1 205 tableaux sur 1 226 tiennent déjà à 9 pt."""
    entetes = ["Poste", "2027", "2028", "2029"]
    lignes = [["Chiffre d'affaires", "40 000 €", "62 000 €", "80 000 €"]]

    assert forme_du_tableau(entetes, lignes) == FORME_TABLEAU_CONFORTABLE


def test_une_adresse_web_se_replie_a_ses_barres_et_ne_reduit_pas_le_tableau() -> None:
    """CONTRE-ÉPREUVE : le moteur passe à la ligne après « / » et « - ».

    Mesurer l'adresse entière comme un seul mot aurait réduit au minimum tous
    les tableaux de sources, qui se lisaient très bien (PDF de `99a142b0`,
    chapitre 21).
    """
    entetes = ["Organisme", "Publication", "Lien"]
    lignes = [[
        "Numeum",
        "Observatoire de conjoncture du marché du numérique",
        "https://numeum.fr/wp-content/uploads/2025/12/"
        "2025-S2-Observatoire-de-conjoncture-NUMEUM-PAC-corrige-1.pdf",
    ]]

    assert forme_du_tableau(entetes, lignes) == FORME_TABLEAU_CONFORTABLE


def test_le_corps_reduit_est_pose_sur_le_texte() -> None:
    """Un corps choisi pour faire tenir les mots, mais pas appliqué, ne sert à rien."""
    entetes = ENTETES_99A
    forme = forme_du_tableau(entetes, LIGNES_99A)
    assert forme != FORME_TABLEAU_CONFORTABLE

    (table,) = _rendre(entetes, LIGNES_99A)

    tailles = {
        run.font.size.pt
        for ligne in table.rows
        for cellule in ligne.cells
        for run in cellule.paragraphs[0].runs
    }
    assert tailles == {forme.taille}


# ── Dernier recours : deux tableaux ──────────────────────────────────────────


ENTETES_LARGES = [
    "Organisation", "Établissement", "Accompagnement", "Positionnement",
    "Représentation", "Communication", "Différenciation", "Responsabilité",
    "Développement",
]
LIGNES_LARGES = [[mot.lower() for mot in ENTETES_LARGES]]


def test_neuf_colonnes_trop_longues_deviennent_deux_tableaux() -> None:
    """Rien ne tient en 17 cm : on coupe le TABLEAU, plus les mots."""
    tables = _rendre(
        ENTETES_LARGES, LIGNES_LARGES, source="Insee, 2026", titre="Grille"
    )

    assert len(tables) == 2
    premieres = [table.rows[0].cells[0].text for table in tables]
    assert premieres == ["Organisation", "Organisation"]
    vues = list(itertools.chain.from_iterable(
        [cellule.text for cellule in table.rows[0].cells[1:]] for table in tables
    ))
    assert vues == ENTETES_LARGES[1:]


def test_apres_scission_le_titre_ouvre_et_la_source_ferme() -> None:
    """Titre au-dessus de la première moitié, « (suite) » sur la seconde,
    source une seule fois, sous la seconde."""
    document = charger_gabarit()
    palette = construire_palette(primaire=REF_PRUNE, fond_clair=REF_CREME)
    corps = document.element.body
    # Le corps finit par `w:sectPr` : ce qu'on ajoute s'insère juste avant.
    depart = len(corps) - 1
    composants.tableau(
        document, palette, ENTETES_LARGES, LIGNES_LARGES,
        source="Insee, 2026", titre="Grille",
    )
    sequence = []
    for element in list(corps)[depart:-1]:
        if element.tag == qn("w:tbl"):
            sequence.append("TABLEAU")
        else:
            texte = "".join(t.text or "" for t in element.iter(qn("w:t"))).strip()
            if texte:
                sequence.append(texte)

    assert sequence[0] == "Grille"
    assert sequence[1] == "TABLEAU"
    assert sequence[2] == "Grille (suite)"
    assert sequence[3] == "TABLEAU"
    assert sum("Insee, 2026" in morceau for morceau in sequence) == 1
    assert "Insee, 2026" in sequence[4]


def test_chaque_moitie_tient_sans_couper_de_mot() -> None:
    for table in _rendre(ENTETES_LARGES, LIGNES_LARGES):
        entetes = [cellule.text for cellule in table.rows[0].cells]
        lignes = [[cellule.text for cellule in ligne.cells] for ligne in table.rows[1:]]
        forme = forme_du_tableau(entetes, lignes)
        planchers = _planchers(entetes, lignes, len(entetes), forme)

        assert all(
            largeur >= plancher
            for largeur, plancher in zip(_grille(table), planchers, strict=True)
        )


def test_le_tableau_du_99a142b0_n_est_pas_coupe_en_deux() -> None:
    """CONTRE-ÉPREUVE : resserrer suffit, on ne coupe que si rien ne tient."""
    assert len(_rendre(ENTETES_99A, LIGNES_99A)) == 1
